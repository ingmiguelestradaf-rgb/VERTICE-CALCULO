"""Costos: Factor de Salario Real, precios unitarios (APU), presupuesto y exportación a Excel.

Lee las bases de la app Precios de Obra MX:
  precios.json    materiales (min/max, unidad, fuente, fecha)
  mano-obra.json  salarios, parámetros FSR 2026, cuadrillas, rendimientos, equipo, factor regional
Por omisión busca una copia local (carpeta hermana ../precios-obra-) y si no la hay las descarga de GitHub.
"""
import json, os, urllib.request
from dataclasses import dataclass, field

URL_BASE = "https://raw.githubusercontent.com/ingmiguelestradaf-rgb/precios-obra-/main/"

def _cargar(nombre, ruta=None):
    candidatos = [ruta] if ruta else []
    aqui = os.path.dirname(os.path.abspath(__file__))
    candidatos += [os.path.join(aqui, "..", "..", "precios-obra-", nombre), os.path.join(aqui, "..", "datos", nombre)]
    for c in candidatos:
        if c and os.path.exists(c):
            with open(c, encoding="utf-8") as f:
                return json.load(f)
    with urllib.request.urlopen(URL_BASE + nombre, timeout=30) as r:
        return json.loads(r.read().decode("utf-8"))

class BaseDatos:
    def __init__(self, ruta_precios=None, ruta_mano_obra=None, factor_regional=None):
        self.precios = {p["_id"]: p for p in _cargar("precios.json", ruta_precios)["precios"]}
        self.mo = _cargar("mano-obra.json", ruta_mano_obra)
        self.factor_regional = self.mo["factor_regional"]["tijuana"] if factor_regional is None else factor_regional
        self.salarios = {c["id"]: round((c["sd_ref_min"] + c["sd_ref_max"]) / 2) for c in self.mo["categorias"]}
        self.cotizaciones = {}  # id -> precio sin IVA puesto en obra (sustituye a la base)

    # --- materiales ---
    def precio_material(self, mid):
        if mid in self.cotizaciones:
            return self.cotizaciones[mid], "cotización"
        p = self.precios[mid]
        v = (p["min"] + p["max"]) / 2
        if p.get("base") == "mx":
            v *= 1 + self.factor_regional / 100
        return v, f'{p.get("src_t") or p.get("src_u", "")} ({p.get("date", "")}) + {self.factor_regional}% Tijuana'

    def material(self, mid):
        return self.precios[mid]

    # --- FSR ---
    def fsr(self, sd, params=None):
        f = dict(self.mo["fsr"]); imss = dict(f["imss"])
        if params:
            for k, v in params.items():
                (imss if k in imss else f)[k] = v
        uma, sm = self.mo["meta"]["uma"], self.mo["meta"]["sm_zlfn"]
        Tp = 365 + f["aguinaldo_dias"] + f["vacaciones_dias"] * f["prima_vacacional"]
        Tl = 365 - f["domingos"] - f["vacaciones_dias"] - f["festivos"] - f["dias_costumbre"] - f["dias_clima"]
        sbc = min(sd * Tp / 365, 25 * uma)
        tabla = imss["ceav_2026"]
        if sd <= sm * 1.0001:
            ceav = tabla[0][1]
        else:
            ceav = next(p for lim, p in tabla[1:] if sbc / uma <= lim)
        pct = (imss["riesgo_trabajo"] + imss["prestaciones_dinero"] + imss["gastos_medicos_pens"] + imss["invalidez_vida"]
               + imss["guarderias"] + imss["retiro"] + ceav + f["infonavit"])
        dia = sbc * pct / 100 + uma * imss["em_cuota_fija"] / 100 + max(0, sbc - 3 * uma) * imss["em_excedente"] / 100
        rem = sd * Tp
        ps = (dia * 365 + rem * f["isn"] / 100) / rem
        F = Tp / Tl * (1 + ps)
        return {"FSR": F, "Tp": Tp, "Tl": Tl, "Ps": ps, "SBC": sbc, "CEAV": ceav, "salario_real": sd * F}

    def salario_real(self, cat):
        return self.fsr(self.salarios[cat])["salario_real"]

    def cuadrilla(self, cid):
        return next(c for c in self.mo["cuadrillas"] if c["id"] == cid)

    def costo_cuadrilla(self, cid):
        return sum(i["cant"] * self.salario_real(i["cat"]) for i in self.cuadrilla(cid)["integrantes"])

    def equipo(self, eid):
        e = next(x for x in self.mo["equipo"] if x["id"] == eid)
        return e, (e["min"] + e["max"]) / 2

    def rendimiento(self, rid):
        return next(r for r in self.mo["rendimientos"] if r["id"] == rid)

@dataclass
class Insumo:
    tipo: str          # Material, Mano de obra, Herramienta, Equipo, Subcontrato
    descripcion: str
    unidad: str
    cantidad: float
    precio: float
    fuente: str = ""

    @property
    def importe(self):
        return self.cantidad * self.precio

@dataclass
class APU:
    clave: str
    concepto: str
    unidad: str
    partida: str
    insumos: list = field(default_factory=list)
    herramienta_menor: float = 0.03

    @property
    def mano_obra(self):
        return sum(i.importe for i in self.insumos if i.tipo == "Mano de obra")

    @property
    def costo_directo(self):
        return sum(i.importe for i in self.insumos) + self.mano_obra * self.herramienta_menor

    def por_tipo(self):
        d = {}
        for i in self.insumos:
            d[i.tipo] = d.get(i.tipo, 0) + i.importe
        d["Herramienta menor"] = self.mano_obra * self.herramienta_menor
        return d

def apu_desde_base(db, rid, clave=None, rendimiento=None, extra=None):
    """APU a partir de un rendimiento de mano-obra.json. `extra`: lista de Insumo adicionales."""
    r = db.rendimiento(rid)
    rend = rendimiento or (r["rend_min"] + r["rend_max"]) / 2
    a = APU(clave or rid.upper(), r["concepto"], r["unidad"], r["partida"])
    for m in r["materiales"]:
        p = db.material(m["id"])
        v, fuente = db.precio_material(m["id"])
        a.insumos.append(Insumo("Material", f'{p["mat"]} {p["spec"]}', p["unit"], m["cant"], v, fuente))
    q = db.cuadrilla(r["cuadrilla"])
    a.insumos.append(Insumo("Mano de obra", f'{q["nombre"]} (rend. {rend:g} {r["unidad"]}/jor)', "jor", 1 / rend,
                            db.costo_cuadrilla(r["cuadrilla"]), "FSR 2026 Tijuana"))
    for e in r["equipo"]:
        eq, v = db.equipo(e["id"])
        a.insumos.append(Insumo("Equipo", eq["nombre"], eq["unidad"], e["cant"], v, eq["nota"]))
    for i in extra or []:
        a.insumos.append(i)
    return a

@dataclass
class Sobrecosto:
    indirectos_oficina: float = 6.0
    indirectos_campo: float = 6.0
    financiamiento: float = 1.0
    utilidad: float = 10.0
    cargos_adicionales: float = 0.0
    iva: float = 8.0

    @property
    def factor(self):
        ind = (self.indirectos_oficina + self.indirectos_campo) / 100
        return (1 + ind) * (1 + self.financiamiento / 100) * (1 + self.utilidad / 100) * (1 + self.cargos_adicionales / 100)

@dataclass
class Concepto:
    apu: APU
    cantidad: float

@dataclass
class Presupuesto:
    proyecto: str
    cliente: str = ""
    ubicacion: str = "Tijuana, B.C."
    area_construida: float = 0.0   # m²
    sobrecosto: Sobrecosto = field(default_factory=Sobrecosto)
    conceptos: list = field(default_factory=list)

    def agregar(self, apu, cantidad):
        self.conceptos.append(Concepto(apu, cantidad))
        return self

    def pu(self, c):
        return c.apu.costo_directo * self.sobrecosto.factor

    @property
    def costo_directo(self):
        return sum(c.apu.costo_directo * c.cantidad for c in self.conceptos)

    @property
    def subtotal(self):
        return sum(self.pu(c) * c.cantidad for c in self.conceptos)

    @property
    def total(self):
        return self.subtotal * (1 + self.sobrecosto.iva / 100)

    def por_partida(self):
        d = {}
        for c in self.conceptos:
            d[c.apu.partida] = d.get(c.apu.partida, 0) + self.pu(c) * c.cantidad
        return d

    def explosion_insumos(self):
        """Suma de insumos de todo el presupuesto: {(tipo, descripcion, unidad): [cantidad, importe]}."""
        d = {}
        for c in self.conceptos:
            for i in c.apu.insumos:
                k = (i.tipo, i.descripcion, i.unidad)
                v = d.setdefault(k, [0.0, 0.0])
                v[0] += i.cantidad * c.cantidad
                v[1] += i.importe * c.cantidad
        return d

    def sensibilidad(self, tipos=("Material", "Mano de obra"), delta=0.10):
        base = self.subtotal
        out = {}
        for t in tipos:
            imp = sum(i.importe * c.cantidad for c in self.conceptos for i in c.apu.insumos if i.tipo == t)
            out[t] = (imp * delta * self.sobrecosto.factor, imp * delta * self.sobrecosto.factor / base)
        return out

    def a_excel(self, ruta):
        """Excel con fórmulas vivas: Resumen, Catálogo, APU, Explosión de insumos y Sobrecosto."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        wb = Workbook()
        H = Font(bold=True, color="FFFFFF"); F = PatternFill("solid", fgColor="1D2328")
        thin = Side(style="thin", color="CDD2D6"); B = Border(bottom=thin)
        money = '"$"#,##0.00'
        def head(ws, row, vals):
            for j, v in enumerate(vals, 1):
                c = ws.cell(row, j, v); c.font = H; c.fill = F; c.alignment = Alignment(wrap_text=True, vertical="center")
        # Sobrecosto
        so = wb.active; so.title = "Sobrecosto"
        s = self.sobrecosto
        filas = [("Indirectos de oficina %", s.indirectos_oficina), ("Indirectos de campo %", s.indirectos_campo),
                 ("Financiamiento %", s.financiamiento), ("Utilidad %", s.utilidad),
                 ("Cargos adicionales %", s.cargos_adicionales), ("IVA %", s.iva)]
        head(so, 1, ["Concepto", "Valor"])
        for i, (k, v) in enumerate(filas, 2):
            so.cell(i, 1, k); so.cell(i, 2, v)
        so.cell(8, 1, "Factor de sobrecosto").font = Font(bold=True)
        so["B8"] = "=(1+(B2+B3)/100)*(1+B4/100)*(1+B5/100)*(1+B6/100)"; so["B8"].number_format = "0.0000"
        so.column_dimensions["A"].width = 32
        # APU
        ap = wb.create_sheet("APU")
        row = 1; cd_cell = {}
        for c in self.conceptos:
            a = c.apu
            ap.cell(row, 1, a.clave).font = Font(bold=True); ap.cell(row, 2, a.concepto).font = Font(bold=True); ap.cell(row, 5, a.unidad)
            row += 1; head(ap, row, ["Tipo", "Insumo", "Unidad", "Cantidad", "Precio", "Importe", "Fuente"]); row += 1
            first = row
            for i in a.insumos:
                ap.cell(row, 1, i.tipo); ap.cell(row, 2, i.descripcion); ap.cell(row, 3, i.unidad)
                ap.cell(row, 4, round(i.cantidad, 6)); ap.cell(row, 5, round(i.precio, 2)).number_format = money
                ap.cell(row, 6, f"=D{row}*E{row}").number_format = money; ap.cell(row, 7, i.fuente)
                row += 1
            last = row - 1
            ap.cell(row, 1, "Herramienta"); ap.cell(row, 2, f"Herramienta menor ({a.herramienta_menor:.0%} de M.O.)")
            ap.cell(row, 3, "%"); ap.cell(row, 4, a.herramienta_menor)
            ap.cell(row, 5, f'=SUMIF(A{first}:A{last},"Mano de obra",F{first}:F{last})').number_format = money
            ap.cell(row, 6, f"=D{row}*E{row}").number_format = money
            row += 1
            ap.cell(row, 5, "Costo directo").font = Font(bold=True)
            ap.cell(row, 6, f"=SUM(F{first}:F{row-1})").number_format = money
            ap.cell(row, 6).font = Font(bold=True)
            cd_cell[a.clave] = f"APU!F{row}"
            row += 2
        for col, w in zip("ABCDEFG", (14, 52, 8, 11, 13, 14, 60)):
            ap.column_dimensions[col].width = w
        # Catálogo
        ca = wb.create_sheet("Catálogo", 0)
        ca["A1"] = f"Presupuesto — {self.proyecto}"; ca["A1"].font = Font(bold=True, size=14)
        ca["A2"] = f"{self.cliente}  ·  {self.ubicacion}  ·  Vértice Ingeniería, Arquitectura y Proyectos"
        head(ca, 4, ["Partida", "Clave", "Concepto", "Unidad", "Cantidad", "Costo directo", "P.U.", "Importe"])
        r0 = 5
        for k, c in enumerate(self.conceptos):
            r = r0 + k
            ca.cell(r, 1, c.apu.partida); ca.cell(r, 2, c.apu.clave); ca.cell(r, 3, c.apu.concepto); ca.cell(r, 4, c.apu.unidad)
            ca.cell(r, 5, c.cantidad)
            ca.cell(r, 6, f"={cd_cell[c.apu.clave]}").number_format = money
            ca.cell(r, 7, f"=ROUND(F{r}*Sobrecosto!$B$8,2)").number_format = money
            ca.cell(r, 8, f"=E{r}*G{r}").number_format = money
        rl = r0 + len(self.conceptos) - 1
        t = rl + 2
        for j, (lab, f) in enumerate([("Subtotal", f"=SUM(H{r0}:H{rl})"), ("IVA", f"=H{t}*Sobrecosto!B7/100"),
                                      ("Total", f"=H{t}+H{t+1}")]):
            ca.cell(t + j, 7, lab).font = Font(bold=True)
            ca.cell(t + j, 8, f).number_format = money; ca.cell(t + j, 8).font = Font(bold=True)
        for col, w in zip("ABCDEFGH", (16, 12, 58, 8, 11, 14, 14, 16)):
            ca.column_dimensions[col].width = w
        ca.freeze_panes = "A5"
        # Resumen
        rs = wb.create_sheet("Resumen", 0)
        rs["A1"] = f"Resumen por partida — {self.proyecto}"; rs["A1"].font = Font(bold=True, size=14)
        head(rs, 3, ["Partida", "Importe sin IVA", "% del total", "Costo por m²"])
        partidas = list(dict.fromkeys(c.apu.partida for c in self.conceptos))
        for k, p in enumerate(partidas):
            r = 4 + k
            rs.cell(r, 1, p)
            rs.cell(r, 2, f"=SUMIF('Catálogo'!A{r0}:A{rl},A{r},'Catálogo'!H{r0}:H{rl})").number_format = money
            rs.cell(r, 3, f"=B{r}/B{4+len(partidas)}").number_format = "0.0%"
            rs.cell(r, 4, f"=IF({self.area_construida or 0}>0,B{r}/{self.area_construida or 1},0)").number_format = money
        tr = 4 + len(partidas)
        rs.cell(tr, 1, "Total sin IVA").font = Font(bold=True)
        rs.cell(tr, 2, f"=SUM(B4:B{tr-1})").number_format = money
        rs.cell(tr, 4, f"=IF({self.area_construida or 0}>0,B{tr}/{self.area_construida or 1},0)").number_format = money
        rs.cell(tr + 1, 1, "Total con IVA").font = Font(bold=True)
        rs.cell(tr + 1, 2, f"=B{tr}*(1+Sobrecosto!B7/100)").number_format = money
        rs.cell(tr + 3, 1, "Para revisión del DRO. Precios de referencia; sustituir por cotizaciones antes de contratar.")
        for col, w in zip("ABCD", (30, 18, 12, 16)):
            rs.column_dimensions[col].width = w
        # Explosión de insumos
        ex = wb.create_sheet("Insumos")
        head(ex, 1, ["Tipo", "Insumo", "Unidad", "Cantidad total", "Importe (costo directo)"])
        for k, ((tipo, desc, u), (q, imp)) in enumerate(sorted(self.explosion_insumos().items()), 2):
            ex.cell(k, 1, tipo); ex.cell(k, 2, desc); ex.cell(k, 3, u); ex.cell(k, 4, round(q, 3))
            ex.cell(k, 5, round(imp, 2)).number_format = money
        for col, w in zip("ABCDE", (14, 60, 8, 15, 20)):
            ex.column_dimensions[col].width = w
        wb.save(ruta)
        return ruta
