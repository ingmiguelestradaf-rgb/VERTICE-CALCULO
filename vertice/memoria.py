"""Memoria de cálculo y dictamen estructural en Word (.docx) con python-docx."""
from datetime import date
from . import LEYENDA

ESTRUCTURA_MEMORIA = [
    "1. Datos generales y normativa aplicada",
    "2. Descripción y estructuración",
    "3. Materiales",
    "4. Análisis de cargas",
    "5. Análisis sísmico y por viento",
    "6. Diseño de elementos",
    "7. Diseño de cimentación",
    "8. Supuestos a validar y recomendaciones",
    "9. Responsiva",
]

ESTRUCTURA_DICTAMEN = [
    "1. Antecedentes y alcance", "2. Documentos revisados e inspección", "3. Descripción de la estructura",
    "4. Estado actual", "5. Revisión estructural", "6. Conclusiones", "7. Recomendaciones", "8. Responsiva",
]

NORMATIVA_BC = [
    ("Ley de Edificaciones del Estado de Baja California y su reglamento", "Estatal"),
    ("Reglamento de Edificación del municipio del predio (Tijuana u otro)", "Municipal"),
    ("Normas Técnicas Complementarias (Criterios y Acciones, Concreto, Mampostería, Cimentaciones, Sismo)", "Referencia adoptada"),
    ("Manual de Diseño de Obras Civiles CFE — Diseño por Sismo (PRODISIS) y Diseño por Viento", "Acciones"),
    ("ACI 318-19 Building Code Requirements for Structural Concrete", "Apoyo"),
    ("NMX-B-506 / ASTM A615 acero de refuerzo; ASTM C39, C31, C94 concreto", "Materiales y pruebas"),
]

class Documento:
    def __init__(self, titulo, proyecto, ubicacion="Tijuana, B.C.", cliente="", tipo="MEMORIA DE CÁLCULO ESTRUCTURAL"):
        from docx import Document
        from docx.shared import Pt, RGBColor, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        self._Pt, self._RGB, self._Cm, self._AL = Pt, RGBColor, Cm, WD_ALIGN_PARAGRAPH
        self.doc = Document()
        st = self.doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
        for s in self.doc.sections:
            s.left_margin = s.right_margin = Cm(2.2); s.top_margin = s.bottom_margin = Cm(2)
            s.header.paragraphs[0].text = f"Vértice Ingeniería, Arquitectura y Proyectos · {proyecto}"
            s.footer.paragraphs[0].text = LEYENDA
        d = self.doc
        p = d.add_paragraph(); r = p.add_run("VÉRTICE"); r.bold = True; r.font.size = Pt(30); r.font.color.rgb = RGBColor(0x1d, 0x23, 0x28)
        d.add_paragraph("Ingeniería, Arquitectura y Proyectos · Tijuana, Baja California")
        d.add_paragraph()
        p = d.add_paragraph(); r = p.add_run(tipo); r.bold = True; r.font.size = Pt(18)
        p = d.add_paragraph(); r = p.add_run(titulo); r.font.size = Pt(14)
        for k, v in (("Proyecto", proyecto), ("Ubicación", ubicacion), ("Propietario", cliente),
                     ("Fecha", date.today().strftime("%d/%m/%Y")), ("Elaboró", "Ing. Miguel Estrada")):
            if v: d.add_paragraph(f"{k}: {v}")
        d.add_paragraph()
        p = d.add_paragraph(); r = p.add_run(LEYENDA.upper()); r.bold = True; r.font.color.rgb = RGBColor(0xb2, 0x3b, 0x2e)
        d.add_page_break()

    def seccion(self, titulo, nivel=1):
        self.doc.add_heading(titulo, level=nivel)

    def parrafo(self, texto, negrita=False):
        p = self.doc.add_paragraph(); r = p.add_run(texto); r.bold = negrita
        return p

    def vinetas(self, items):
        for it in items:
            self.doc.add_paragraph(str(it), style="List Bullet")

    def formula(self, texto, ref=""):
        p = self.doc.add_paragraph(); r = p.add_run(texto); r.font.name = "Consolas"
        if ref:
            r2 = p.add_run(f"    [{ref}]"); r2.italic = True; r2.font.size = self._Pt(8.5)

    def tabla(self, encabezados, filas, anchos_cm=None):
        t = self.doc.add_table(rows=1, cols=len(encabezados)); t.style = "Light Grid Accent 1"
        for c, h in zip(t.rows[0].cells, encabezados):
            c.text = str(h)
            for r in c.paragraphs[0].runs: r.bold = True
        for f in filas:
            cells = t.add_row().cells
            for c, v in zip(cells, f):
                c.text = _fmt(v)
        if anchos_cm:
            for row in t.rows:
                for c, w in zip(row.cells, anchos_cm): c.width = self._Cm(w)
        self.doc.add_paragraph()
        return t

    def imagen(self, ruta, ancho_cm=16, pie=None):
        self.doc.add_picture(ruta, width=self._Cm(ancho_cm))
        if pie:
            p = self.doc.add_paragraph(pie); p.runs[0].italic = True

    def supuestos(self, filas):
        """filas: (supuesto, valor usado, impacto si cambia)."""
        self.tabla(["Supuesto a validar", "Valor usado", "Impacto si cambia"], filas, [6, 4, 6])

    def responsiva(self):
        self.parrafo("El presente documento se elaboró con base en la información disponible y los supuestos señalados. "
                     "Su validez queda sujeta a la revisión, aprobación y firma del Director Responsable de Obra y, en su caso, "
                     "del Corresponsable en Seguridad Estructural, así como a la autorización de la autoridad municipal.")
        self.doc.add_paragraph("\n\n______________________________\nIng. Miguel Estrada\nDRO / Perito en edificación")

    def guardar(self, ruta):
        self.doc.save(ruta)
        return ruta

def _fmt(v):
    if isinstance(v, bool):
        return "Cumple" if v else "NO CUMPLE"
    if isinstance(v, float):
        return f"{v:,.2f}"
    return str(v)
