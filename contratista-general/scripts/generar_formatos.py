#!/usr/bin/env python3
"""Formatos de control del contratista general (Excel con fórmulas vivas, membrete Vértice).

Uso:
    python generar_formatos.py datos.json carpeta_salida [flujo estimaciones obra personal maquinaria cotizacion]
Sin nombres genera los seis. Después recalcula cada archivo con recalc.py del skill xlsx.

Convención de color (igual en todos los libros):
    celda amarilla con texto azul = dato que se captura
    texto negro                   = fórmula (no se escribe encima)
Parámetros tomados de VERTICE-CALCULO (vertice/costos.py y vertice/programa.py): factor de sobrecosto,
esquema de anticipo/estimaciones/retención y de egresos, salario mínimo ZLFN y FSR 2026.
"""
import json, os, sys
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.chart import LineChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment

OSCURO, ACENTO, AZULV = "1D2328", "B23B2E", "1F5F8B"
F = "Arial"
H_FONT = Font(name=F, bold=True, color="FFFFFF", size=9)
H_FILL = PatternFill("solid", fgColor=OSCURO)
IN_FONT = Font(name=F, color="0000FF", size=9)
IN_FILL = PatternFill("solid", fgColor="FFF2CC")
N_FONT = Font(name=F, size=9)
B_FONT = Font(name=F, size=9, bold=True)
TOT_FILL = PatternFill("solid", fgColor="E3E6E8")
LINEA = Side(style="thin", color="BFC5CA")
BORDE = Border(left=LINEA, right=LINEA, top=LINEA, bottom=LINEA)
MON = '"$"#,##0.00;("$"#,##0.00);"-"'
MON0 = '"$"#,##0;("$"#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
NUM = '#,##0.00;(#,##0.00);"-"'
FECHA = "dd/mm/yyyy"
ROJO = PatternFill("solid", fgColor="F8D7D3")
AMBAR = PatternFill("solid", fgColor="FCE8C3")
VERDE = PatternFill("solid", fgColor="D9EDDF")

def g(D, ruta, defecto=""):
    o = D
    for k in ruta.split("."):
        o = o.get(k) if isinstance(o, dict) else None
        if o is None:
            return defecto
    return o if o != "" else defecto

# ---------------- ayudantes de hoja ----------------
def membrete(ws, titulo, D, ultima_col):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ultima_col)
    c = ws.cell(1, 1, "VÉRTICE  ·  Ingeniería, Arquitectura y Proyectos")
    c.font = Font(name=F, bold=True, size=13, color=OSCURO)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ultima_col)
    c = ws.cell(2, 1, titulo.upper()); c.font = Font(name=F, bold=True, size=11, color=ACENTO)
    ws.cell(3, 1, "Obra:").font = B_FONT
    ws.cell(3, 2, g(D, "obra.nombre")).font = N_FONT
    ws.cell(4, 1, "Contrato:").font = B_FONT
    ws.cell(4, 2, g(D, "contrato.numero")).font = N_FONT
    for col in range(1, ultima_col + 1):
        ws.cell(2, col).border = Border(bottom=Side(style="medium", color=ACENTO))

def encabezados(ws, fila, cols):
    for j, (nombre, ancho) in enumerate(cols, 1):
        c = ws.cell(fila, j, nombre); c.font = H_FONT; c.fill = H_FILL; c.border = BORDE
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if ancho:
            ws.column_dimensions[L(j)].width = ancho
    ws.row_dimensions[fila].height = 30

def celda(ws, r, c, v=None, fmt=None, entrada=False, negrita=False):
    x = ws.cell(r, c, v)
    x.font = IN_FONT if entrada else (B_FONT if negrita else N_FONT)
    if entrada:
        x.fill = IN_FILL
    if fmt:
        x.number_format = fmt
    x.border = BORDE
    return x

def fila_total(ws, r, c1, c2):
    for c in range(c1, c2 + 1):
        x = ws.cell(r, c); x.fill = TOT_FILL; x.font = B_FONT; x.border = BORDE

def param(ws, r, etiqueta, valor, fmt=None, entrada=True, nota=None):
    a = ws.cell(r, 1, etiqueta); a.font = N_FONT; a.border = BORDE
    b = celda(ws, r, 2, valor, fmt, entrada=entrada)
    if nota:
        n = ws.cell(r, 3, nota); n.font = Font(name=F, size=8, italic=True, color="5B656E")
    return b

def instrucciones(wb, titulo, lineas):
    ws = wb.active; ws.title = "Instrucciones"; ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 110
    ws["A1"] = "VÉRTICE  ·  Ingeniería, Arquitectura y Proyectos"; ws["A1"].font = Font(name=F, bold=True, size=13, color=OSCURO)
    ws["A2"] = titulo.upper(); ws["A2"].font = Font(name=F, bold=True, size=11, color=ACENTO)
    ws["A4"] = "Captura"; ws["A4"].font = IN_FONT; ws["A4"].fill = IN_FILL
    ws["A5"] = "Celdas amarillas con texto azul = datos que capturas. Texto negro = fórmulas: no escribas encima. La fila de ejemplo muestra el formato esperado; sustitúyela por tus datos."
    ws["A5"].font = N_FONT; ws["A5"].alignment = Alignment(wrap_text=True)
    r = 7
    for ln in lineas:
        c = ws.cell(r, 1, ln); c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = B_FONT if ln.endswith(":") else N_FONT
        r += 1
    return ws

def lista(ws, rango, opciones):
    dv = DataValidation(type="list", formula1='"' + ",".join(opciones) + '"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(rango)

def firmas(ws, fila, etiquetas, col_ini=1, salto=3):
    for i, et in enumerate(etiquetas):
        c = col_ini + i * salto
        ws.cell(fila, c, "_" * 28).font = N_FONT
        ws.cell(fila + 1, c, et).font = B_FONT

def fecha_inicio(D):
    t = g(D, "contrato.fecha_inicio")
    try:
        return date.fromisoformat(t)
    except Exception:
        hoy = date.today()
        return hoy - timedelta(days=hoy.weekday())

def guardar(wb, carpeta, nombre):
    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
        if ws.title not in ("Instrucciones", "Cotización"):
            ws.page_setup.orientation = "landscape"
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
    ruta = os.path.join(carpeta, nombre)
    wb.save(ruta)
    print(ruta)
    return ruta

# ======================================================================
# 1. FLUJO DE CAJA
# ======================================================================
def flujo(D, carpeta):
    wb = Workbook()
    instrucciones(wb, "Flujo de caja semanal de obra", [
        "Para qué sirve:",
        "Proyecta cobros y pagos semana a semana a partir del costo programado, el anticipo, las estimaciones, la retención del fondo de garantía y el crédito de proveedores; calcula cuánto dinero tienes que poner (necesidad máxima de financiamiento) y lo que cuesta.",
        "Cómo se llena:",
        "1. En «Datos» captura el monto del contrato, el costo presupuestado y las condiciones de pago del contrato (mismas cláusulas de ANTICIPO, ESTIMACIONES y FONDO DE GARANTÍA).",
        "2. En «Flujo semanal», columna C, captura el costo programado de cada semana (sale del programa de obra de VERTICE-CALCULO: programa.costo_por_semana()). La celda de control debajo de la tabla debe coincidir con el costo presupuestado.",
        "3. Cada semana captura en Q y R lo que realmente cobraste y pagaste: el libro compara el acumulado real contra el proyectado.",
        "Notas:",
        "Todo está SIN IVA: el IVA se traslada y acredita aparte. Para el IVA en la frontera norte usa la tasa del Decreto vigente (8 % hasta el 31-dic-2026).",
        "La lógica es la misma de vertice/programa.py → flujo_de_caja(): anticipo en la semana 1, estimaciones cada N semanas, cobro con desfase, material con crédito, nómina la misma semana y el resto con una semana de desfase.",
    ])
    d = wb.create_sheet("Datos"); membrete(d, "Flujo de caja — datos", D, 3)
    d.column_dimensions["A"].width = 46; d.column_dimensions["B"].width = 18; d.column_dimensions["C"].width = 60
    monto = float(g(D, "contrato.monto", 1000000))
    param(d, 6, "Fecha de inicio de obra", fecha_inicio(D), FECHA)
    param(d, 7, "Monto del contrato (sin IVA)", monto, MON)
    param(d, 8, "Costo total presupuestado (sin IVA)", round(monto / 1.25, 2), MON, nota="Costo directo + indirectos de campo, sin utilidad (de tu presupuesto)")
    param(d, 9, "Anticipo", float(g(D, "contrato.anticipo_pct", 30)) / 100, PCT)
    param(d, 10, "Fondo de garantía retenido", float(g(D, "contrato.fondo_garantia_pct", 5)) / 100, PCT)
    param(d, 11, "Periodo de estimación (semanas)", 2, "0")
    param(d, 12, "Semanas que tarda el cliente en pagar", 2, "0")
    param(d, 13, "Semanas después del fin para devolver el fondo", 4, "0")
    param(d, 14, "Fracción del costo que es material", 0.55, PCT)
    param(d, 15, "Crédito de proveedores (semanas)", 2, "0")
    param(d, 16, "Fracción del costo que es mano de obra", 0.35, PCT, nota="Se paga la misma semana (nómina)")
    param(d, 17, "Resto (subcontratos e indirectos), 1 semana de desfase", "=1-B14-B16", PCT, entrada=False)
    param(d, 18, "Tasa anual de financiamiento", 0.14, PCT)
    param(d, 19, "IVA (sólo informativo)", float(g(D, "contrato.iva_pct", 8)) / 100, PCT)
    d.cell(21, 1, "RESULTADOS").font = Font(name=F, bold=True, color=ACENTO)
    param(d, 22, "Semana final de obra", "=SUMPRODUCT(MAX(('Flujo semanal'!C8:C59>0)*'Flujo semanal'!A8:A59))", "0", entrada=False)
    param(d, 23, "Necesidad máxima de financiamiento", "=MAX(0,-MIN('Flujo semanal'!O8:O59))", MON, entrada=False)
    param(d, 24, "Semana de mayor necesidad", "=IF(B23=0,\"-\",INDEX('Flujo semanal'!A8:A59,MATCH(MIN('Flujo semanal'!O8:O59),'Flujo semanal'!O8:O59,0)))", "0", entrada=False)
    param(d, 25, "Costo financiero estimado", "=SUM('Flujo semanal'!P8:P59)", MON, entrada=False)
    param(d, 26, "Utilidad bruta proyectada (antes de financiamiento)", "='Flujo semanal'!O59", MON, entrada=False)
    param(d, 27, "Utilidad neta proyectada", "=B26-B25", MON, entrada=False)
    param(d, 28, "Margen sobre contrato", "=IF(B7=0,0,B27/B7)", PCT, entrada=False)
    param(d, 29, "Última semana con captura real", "=IFERROR(LOOKUP(2,1/('Flujo semanal'!S8:S59<>\"\"),'Flujo semanal'!A8:A59),\"-\")", "0", entrada=False)
    param(d, 30, "Diferencia real vs proyectado a esa semana", "=IFERROR(LOOKUP(2,1/('Flujo semanal'!T8:T59<>\"\"),'Flujo semanal'!T8:T59),0)", MON, entrada=False, nota="Negativo = traes menos caja de la proyectada")

    f = wb.create_sheet("Flujo semanal"); membrete(f, "Flujo de caja semanal (sin IVA)", D, 20)
    cols = [("Semana", 7), ("Inicia", 11), ("Costo programado", 14), ("% avance semanal", 9), ("Estimación", 14),
            ("Amortiz. anticipo", 13), ("Fondo garantía", 12), ("Neto a cobrar", 14), ("COBROS proyectados", 14),
            ("Pago material", 13), ("Pago mano de obra", 13), ("Otros pagos", 12), ("PAGOS proyectados", 14),
            ("Saldo semanal", 14), ("Acumulado proyectado", 15), ("Costo financiero", 11),
            ("Cobro REAL", 13), ("Pago REAL", 13), ("Acumulado REAL", 15), ("Real − proyectado", 14)]
    encabezados(f, 7, cols)
    costo = monto / 1.25
    curva = [0.03, 0.07, 0.11, 0.14, 0.15, 0.15, 0.13, 0.10, 0.07, 0.05]  # curva S de ejemplo (10 semanas)
    R0, R1 = 8, 59
    A = f"$A${R0}:$A${R1}"; Dc = f"$D${R0}:$D${R1}"
    for i, r in enumerate(range(R0, R1 + 1), 1):
        celda(f, r, 1, i, "0")
        celda(f, r, 2, f"=Datos!$B$6+7*(A{r}-1)", FECHA)
        celda(f, r, 3, round(costo * curva[i - 1], 2) if i <= len(curva) else None, MON0, entrada=True)
        celda(f, r, 4, f"=IF(Datos!$B$8=0,0,C{r}/Datos!$B$8)", PCT)
        n, sf = "Datos!$B$11", "Datos!$B$22"
        celda(f, r, 5, f"=IF(AND(A{r}<={sf},OR(MOD(A{r},{n})=0,A{r}={sf})),SUMIFS({Dc},{A},\">\"&(A{r}-IF(MOD(A{r},{n})=0,{n},MOD(A{r},{n}))),{A},\"<=\"&A{r})*Datos!$B$7,0)", MON0)
        celda(f, r, 6, f"=E{r}*Datos!$B$9", MON0)
        celda(f, r, 7, f"=E{r}*Datos!$B$10", MON0)
        celda(f, r, 8, f"=E{r}-F{r}-G{r}", MON0)
        celda(f, r, 9, f"=IF(A{r}=1,Datos!$B$7*Datos!$B$9,0)+IF(A{r}-Datos!$B$12>=1,INDEX($H${R0}:$H${R1},A{r}-Datos!$B$12),0)"
                       f"+IF(A{r}={sf}+Datos!$B$12+Datos!$B$13,SUM($G${R0}:$G${R1}),0)", MON0)
        celda(f, r, 10, f"=IF(A{r}-Datos!$B$15>=1,INDEX($C${R0}:$C${R1},A{r}-Datos!$B$15),0)*Datos!$B$14", MON0)
        celda(f, r, 11, f"=C{r}*Datos!$B$16", MON0)
        celda(f, r, 12, f"=IF(A{r}>1,INDEX($C${R0}:$C${R1},A{r}-1),0)*Datos!$B$17", MON0)
        celda(f, r, 13, f"=J{r}+K{r}+L{r}", MON0)
        celda(f, r, 14, f"=I{r}-M{r}", MON0)
        celda(f, r, 15, f"=N{r}" if r == R0 else f"=O{r-1}+N{r}", MON0)
        celda(f, r, 16, f"=MAX(0,-O{r})*((1+Datos!$B$18)^(1/52)-1)", MON0)
        celda(f, r, 17, None, MON0, entrada=True)
        celda(f, r, 18, None, MON0, entrada=True)
        celda(f, r, 19, f"=IF(AND(Q{r}=\"\",R{r}=\"\"),\"\",SUM($Q${R0}:Q{r})-SUM($R${R0}:R{r}))", MON0)
        celda(f, r, 20, f"=IF(S{r}=\"\",\"\",S{r}-O{r})", MON0)
    tr = R1 + 1
    f.cell(tr, 1, "TOTAL")
    for c in [3, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 16, 17, 18]:
        f.cell(tr, c, f"=SUM({L(c)}{R0}:{L(c)}{R1})").number_format = MON0
    fila_total(f, tr, 1, 20)
    f.cell(tr + 1, 2, "Control:").font = B_FONT
    f.cell(tr + 1, 3, f"=IF(ABS(C{tr}-Datos!B8)<1,\"Cuadra con el presupuesto\",\"No cuadra: revisa C\")").font = Font(name=F, size=9, bold=True, color=ACENTO)
    f.cell(tr + 2, 2, "Columnas Q y R: captura lo cobrado y pagado cada semana; S y T se llenan solas y comparan contra lo proyectado.").font = Font(name=F, size=8, italic=True)
    f.freeze_panes = "C8"
    ch = LineChart(); ch.title = "Saldo acumulado proyectado"; ch.height = 8; ch.width = 22
    ch.y_axis.title = "Pesos (sin IVA)"; ch.x_axis.title = "Semana"
    ch.add_data(Reference(f, min_col=15, min_row=7, max_row=R1), titles_from_data=True)
    ch.set_categories(Reference(f, min_col=1, min_row=R0, max_row=R1))
    d.add_chart(ch, "E6")
    return guardar(wb, carpeta, "Flujo_de_Caja.xlsx")

# ======================================================================
# 2. ESTIMACIONES
# ======================================================================
CONCEPTOS_EJ = [  # clave, concepto, unidad, cantidad, PU (sin IVA)
    ("PRE-01", "Demolición de muros de block y retiro de escombro fuera de obra", "m²", 45, 185.0),
    ("TAB-01", "Muro de tablaroca 1 cara, bastidor cal. 26 @ 61 cm, juntas y cinta", "m²", 210, 520.0),
    ("TAB-02", "Plafón de tablaroca a 3.00 m, incluye colgantes y canaleta", "m²", 180, 610.0),
    ("ELE-01", "Salida eléctrica de contacto dúplex, cable cal. 12, tubería conduit", "sal", 48, 890.0),
    ("ELE-02", "Salida de alumbrado con luminaria LED de empotrar 60×60", "sal", 36, 1650.0),
    ("HID-01", "Salida hidráulica y sanitaria para lavabo, PVC y CPVC", "sal", 6, 2350.0),
    ("PIS-01", "Piso porcelánico 60×60 asentado con adhesivo, incluye boquilla", "m²", 175, 735.0),
    ("PIN-01", "Pintura vinílica 2 manos sobre tablaroca, incluye sellador", "m²", 560, 98.0),
]

def estimaciones(D, carpeta):
    wb = Workbook()
    instrucciones(wb, "Estimación de obra con números generadores", [
        "Para qué sirve:",
        "Integra la estimación (cobro periódico) de un contrato a precios unitarios: catálogo contratado, números generadores, acumulados, carátula con amortización de anticipo, fondo de garantía, IVA y neto a pagar, y reporte fotográfico.",
        "Cómo se llena cada periodo:",
        "1. «Datos»: número de estimación, periodo, porcentajes del contrato (cláusulas de ANTICIPO y ESTIMACIONES).",
        "2. «Catálogo»: el catálogo del Anexo B del contrato (clave, concepto, unidad, cantidad y PU). La «cantidad acumulada anterior» es la acumulada de la estimación pasada: copia la columna I de la estimación anterior como valores.",
        "3. «Generadores»: una fila por medición (eje, nivel, largo, ancho, alto, piezas). Las celdas vacías no cuentan en el producto. Usa signo -1 para descontar huecos. La cantidad de cada clave pasa sola al catálogo.",
        "4. «Carátula»: revisa y firma. Si el contrato prevé la retención de IVA del art. 1-A fr. IV LIVA, captura 6 % en Datos; si no, déjalo en 0.",
        "Reglas que ya trae el libro:",
        "Marca EXCEDE cuando el acumulado supera la cantidad contratada: eso es un volumen adicional que requiere orden de cambio o PU extraordinario (cláusula CONCEPTOS FUERA DE CATÁLOGO).",
        "La base del IVA es la estimación menos la amortización del anticipo (el anticipo ya se facturó con IVA). El fondo de garantía es una retención de pago y no reduce la base del IVA.",
        "Para medir cantidades en campo apóyate en las calculadoras de ObraLab (concreto, block, tablaroca, plafón, pintura, acero).",
    ])
    d = wb.create_sheet("Datos"); membrete(d, "Estimación — datos del contrato", D, 3)
    d.column_dimensions["A"].width = 40; d.column_dimensions["B"].width = 22; d.column_dimensions["C"].width = 50
    ini = fecha_inicio(D)
    param(d, 6, "Contratista", g(D, "empresa.nombre_comercial"))
    param(d, 7, "Cliente / propietario", g(D, "contraparte.nombre"))
    param(d, 8, "Estimación número", 1, "0")
    param(d, 9, "Periodo del", ini, FECHA)
    param(d, 10, "Periodo al", ini + timedelta(days=13), FECHA)
    param(d, 11, "Monto del contrato (sin IVA)", "=Catálogo!F49", MON, entrada=False, nota="Suma del catálogo contratado")
    param(d, 12, "Anticipo otorgado", float(g(D, "contrato.anticipo_pct", 30)) / 100, PCT)
    param(d, 13, "Fondo de garantía", float(g(D, "contrato.fondo_garantia_pct", 5)) / 100, PCT)
    param(d, 14, "IVA", float(g(D, "contrato.iva_pct", 8)) / 100, PCT, nota="8 % región fronteriza norte (Decreto vigente hasta 31-dic-2026)")
    param(d, 15, "Retención de IVA (si aplica)", 0.06 if g(D, "subcontrato.retencion_iva_aplica", False) else 0, PCT, nota="Sólo si aplica art. 1-A fr. IV LIVA; normalmente 0")
    param(d, 16, "Otras deducciones (penas, cargos)", 0, MON)
    param(d, 17, "Concepto de otras deducciones", "")

    c = wb.create_sheet("Catálogo"); membrete(c, "Catálogo, avance y estimación del periodo", D, 14)
    cols = [("Clave", 9), ("Concepto", 46), ("Unidad", 7), ("Cantidad contratada", 11), ("P.U. (sin IVA)", 12), ("Importe contratado", 14),
            ("Cant. acumulada anterior", 11), ("Cant. esta estimación", 11), ("Cant. acumulada actual", 11), ("Importe esta estimación", 14),
            ("Importe acumulado", 14), ("Cant. por ejecutar", 11), ("% avance", 8), ("Alerta", 10)]
    encabezados(c, 7, cols)
    R0, R1 = 8, 48
    for i, r in enumerate(range(R0, R1 + 1)):
        ej = CONCEPTOS_EJ[i] if i < len(CONCEPTOS_EJ) else None
        celda(c, r, 1, ej[0] if ej else None, entrada=True)
        x = celda(c, r, 2, ej[1] if ej else None, entrada=True); x.alignment = Alignment(wrap_text=True)
        celda(c, r, 3, ej[2] if ej else None, entrada=True)
        celda(c, r, 4, ej[3] if ej else None, NUM, entrada=True)
        celda(c, r, 5, ej[4] if ej else None, MON, entrada=True)
        celda(c, r, 6, f"=D{r}*E{r}", MON)
        celda(c, r, 7, (round(ej[3] * 0.3, 2) if ej and i < 4 else 0) if ej else None, NUM, entrada=True)
        celda(c, r, 8, f"=IF(A{r}=\"\",0,SUMIF(Generadores!$A$8:$A$307,A{r},Generadores!$J$8:$J$307))", NUM)
        celda(c, r, 9, f"=G{r}+H{r}", NUM)
        celda(c, r, 10, f"=H{r}*E{r}", MON)
        celda(c, r, 11, f"=I{r}*E{r}", MON)
        celda(c, r, 12, f"=D{r}-I{r}", NUM)
        celda(c, r, 13, f"=IF(D{r}=0,0,I{r}/D{r})", PCT)
        celda(c, r, 14, f"=IF(AND(D{r}>0,I{r}>D{r}+0.0001),\"EXCEDE\",\"\")")
    tr = R1 + 1
    c.cell(tr, 2, "TOTALES")
    for col in (6, 10, 11):
        c.cell(tr, col, f"=SUM({L(col)}{R0}:{L(col)}{R1})").number_format = MON
    c.cell(tr, 13, f"=IF(F{tr}=0,0,K{tr}/F{tr})").number_format = PCT
    fila_total(c, tr, 1, 14)
    c.conditional_formatting.add(f"N{R0}:N{R1}", FormulaRule(formula=[f'N{R0}="EXCEDE"'], fill=ROJO))
    c.freeze_panes = "C8"

    gsh = wb.create_sheet("Generadores"); membrete(gsh, "Números generadores", D, 11)
    cols = [("Clave", 9), ("Eje / nivel / ubicación", 22), ("Descripción de la medición", 34), ("Largo", 9), ("Ancho", 9), ("Alto / espesor", 9),
            ("Piezas", 8), ("Signo (1 / -1)", 8), ("Parcial", 10), ("Cantidad", 10), ("Croquis / foto", 14)]
    encabezados(gsh, 7, cols)
    ejemplos = [("PRE-01", "Eje 1 / A-C", "Muro divisorio a demoler", 9.0, None, 2.8, 1, 1),
                ("PRE-01", "Eje 1 / A-C", "Hueco de puerta existente", 0.9, None, 2.1, 1, -1),
                ("TAB-01", "Ejes B-D / 2", "Muro de tablaroca sala de juntas", 12.5, None, 3.0, 1, 1),
                ("TAB-02", "Área de ventas", "Plafón zona A", 8.0, 7.5, None, 1, 1),
                ("ELE-01", "Muro eje C", "Contactos dúplex", None, None, None, 12, 1),
                ("PIS-01", "Área de ventas", "Piso porcelánico zona A", 8.0, 7.5, None, 1, 1)]
    for i, r in enumerate(range(8, 308)):
        ej = ejemplos[i] if i < len(ejemplos) else None
        for j in range(8):
            v = ej[j] if ej else None
            celda(gsh, r, j + 1, v, NUM if 3 <= j <= 6 else ("0" if j == 7 else None), entrada=True)
        celda(gsh, r, 9, f"=IF(COUNT(D{r}:G{r})=0,0,PRODUCT(D{r}:G{r}))", NUM)
        celda(gsh, r, 10, f"=I{r}*IF(H{r}=\"\",1,H{r})", NUM)
        celda(gsh, r, 11, None, entrada=True)
    lista(gsh, "H8:H307", ["1", "-1"])
    gsh.freeze_panes = "D8"

    k = wb.create_sheet("Carátula"); membrete(k, "Carátula de estimación", D, 6)
    for col, w in zip("ABCDEF", (44, 20, 4, 36, 20, 4)):
        k.column_dimensions[col].width = w
    k.cell(6, 1, "Estimación No."); k.cell(6, 2, "=Datos!B8")
    k.cell(7, 1, "Periodo"); k.cell(7, 2, "=TEXT(Datos!B9,\"dd/mm/yyyy\")&\" al \"&TEXT(Datos!B10,\"dd/mm/yyyy\")")
    k.cell(8, 1, "Contratista"); k.cell(8, 2, "=Datos!B6")
    k.cell(9, 1, "Cliente"); k.cell(9, 2, "=Datos!B7")
    for r in range(6, 10):
        k.cell(r, 1).font = B_FONT; k.cell(r, 2).font = N_FONT
    filas = [
        ("Importe de trabajos ejecutados en el periodo", "=Catálogo!J49"),
        ("(-) Amortización del anticipo", "=-ROUND(B12*Datos!B12,2)"),
        ("Subtotal (base del IVA)", "=B12+B13"),
        ("(+) IVA", "=ROUND(B14*Datos!B14,2)"),
        ("Total de la estimación (importe del CFDI)", "=B14+B15"),
        ("(-) Retención de IVA", "=-ROUND(B14*Datos!B15,2)"),
        ("(-) Retención fondo de garantía", "=-ROUND(B12*Datos!B13,2)"),
        ("(-) Otras deducciones", "=-Datos!B16"),
        ("NETO A PAGAR", "=SUM(B16:B19)"),
    ]
    for i, (et, fo) in enumerate(filas, 12):
        fnt = B_FONT if et.startswith(("Subtotal", "Total", "NETO")) else N_FONT
        a = k.cell(i, 1, et); a.font = fnt; a.border = BORDE
        b = k.cell(i, 2, fo); b.number_format = MON; b.font = fnt; b.border = BORDE
    fila_total(k, 20, 1, 2)
    res = [
        ("Monto del contrato", "=Datos!B11"),
        ("Estimado acumulado (incluye ésta)", "=Catálogo!K49"),
        ("Avance financiero", "=IF(E12=0,0,E13/E12)"),
        ("Anticipo otorgado", "=ROUND(E12*Datos!B12,2)"),
        ("Anticipo amortizado acumulado", "=ROUND(E13*Datos!B12,2)"),
        ("Saldo de anticipo por amortizar", "=E15-E16"),
        ("Fondo de garantía acumulado", "=ROUND(E13*Datos!B13,2)"),
        ("Saldo del contrato por estimar", "=E12-E13"),
    ]
    k.cell(11, 4, "RESUMEN DEL CONTRATO").font = Font(name=F, bold=True, color=ACENTO)
    for i, (et, fo) in enumerate(res, 12):
        a = k.cell(i, 4, et); a.font = N_FONT; a.border = BORDE
        b = k.cell(i, 5, fo); b.number_format = PCT if "Avance" in et else MON; b.font = N_FONT; b.border = BORDE
    k.cell(22, 1, "Alertas del catálogo:").font = B_FONT
    k.cell(22, 2, "=IF(COUNTIF(Catálogo!N8:N48,\"EXCEDE\")=0,\"Sin conceptos excedidos\",COUNTIF(Catálogo!N8:N48,\"EXCEDE\")&\" concepto(s) exceden lo contratado: requieren orden de cambio\")").font = Font(name=F, size=9, bold=True, color=ACENTO)
    k.cell(23, 1, "El pago de esta estimación no implica la aceptación de los trabajos (ver cláusula de pago del contrato).").font = Font(name=F, size=8, italic=True)
    firmas(k, 27, ["Elaboró — superintendente del contratista", "Revisó — supervisión", "Autorizó — propietario"], 1, 2)
    k.column_dimensions["C"].width = 4

    fo = wb.create_sheet("Reporte fotográfico"); membrete(fo, "Reporte fotográfico de la estimación", D, 4)
    for col in "ABCD":
        fo.column_dimensions[col].width = 30
    r = 6
    for n in range(1, 9):
        cc = 1 + ((n - 1) % 2) * 2
        if n % 2 == 1 and n > 1:
            r += 16
        fo.merge_cells(start_row=r, start_column=cc, end_row=r + 12, end_column=cc + 1)
        x = fo.cell(r, cc, f"Foto {n}\n(inserta la imagen aquí)"); x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        x.font = Font(name=F, size=9, color="9AA3AA")
        for rr in range(r, r + 13):
            for c2 in (cc, cc + 1):
                fo.cell(rr, c2).border = BORDE
        celda(fo, r + 13, cc, "Clave / ubicación:", entrada=False)
        celda(fo, r + 13, cc + 1, None, entrada=True)
        celda(fo, r + 14, cc, "Descripción:", entrada=False)
        celda(fo, r + 14, cc + 1, None, entrada=True)
    wb.move_sheet("Carátula", offset=-3)
    return guardar(wb, carpeta, "Estimaciones.xlsx")

# ======================================================================
# 3. CONTROL DE OBRA
# ======================================================================
def control_obra(D, carpeta):
    wb = Workbook()
    instrucciones(wb, "Control de obra", [
        "Para qué sirve:",
        "Lleva el avance físico-financiero con curva S y valor ganado, el reporte diario, las órdenes de cambio, los pendientes y RFI, y las liberaciones y pruebas de calidad de la obra.",
        "Hojas:",
        "Avance: captura el importe de cada partida y su % de avance real; abajo, cada semana, el % programado (de tu programa) y el real. Calcula SPI y CPI como la calculadora de valor ganado de ObraLab (SPI = EV/PV, CPI = EV/AC).",
        "Reporte diario: una fila por día. Es tu respaldo para la bitácora oficial que exige el Reglamento de la Ley de Edificaciones para el Municipio de Tijuana (art. 17); anota el número de asiento.",
        "Órdenes de cambio: ningún trabajo extra se ejecuta sin orden firmada (cláusula de ÓRDENES DE CAMBIO). El total autorizado se suma al contrato.",
        "Pendientes y RFI: el estatus y los días vencidos se calculan solos contra la fecha de hoy.",
        "Liberaciones y pruebas: para concreto captura f'c de proyecto y resistencia obtenida; marca automáticamente si cumple. Para la liberación de colados usa también los checklists de ObraLab (Liberación de colado, Cimentación, Estructura).",
    ])
    a = wb.create_sheet("Avance"); membrete(a, "Avance físico-financiero y valor ganado", D, 11)
    encabezados(a, 6, [("Partida", 34), ("Importe contratado", 15), ("Peso", 8), ("% avance real", 10), ("Avance ponderado", 10),
                       ("Importe ejecutado", 15)])
    for col, w in zip("GHIJK", (3, 14, 14, 10, 10)):
        a.column_dimensions[col].width = w
    partidas = [("Preliminares y demoliciones", 0.04, 1.0), ("Albañilería y tablaroca", 0.22, 0.65), ("Plafones", 0.14, 0.40),
                ("Instalación eléctrica", 0.20, 0.35), ("Instalación hidrosanitaria", 0.06, 0.50), ("Pisos y recubrimientos", 0.16, 0.0),
                ("Pintura y acabados", 0.10, 0.0), ("Limpieza y entrega", 0.08, 0.0)]
    monto = float(g(D, "contrato.monto", 1000000))
    for i in range(15):
        r = 7 + i
        p = partidas[i] if i < len(partidas) else None
        celda(a, r, 1, p[0] if p else None, entrada=True)
        celda(a, r, 2, round(monto * p[1], 2) if p else None, MON0, entrada=True)
        celda(a, r, 3, f"=IF($B$22=0,0,B{r}/$B$22)", PCT)
        celda(a, r, 4, p[2] if p else None, PCT, entrada=True)
        celda(a, r, 5, f"=C{r}*D{r}", PCT)
        celda(a, r, 6, f"=B{r}*D{r}", MON0)
    a.cell(22, 1, "TOTAL")
    a.cell(22, 2, "=SUM(B7:B21)").number_format = MON0
    a.cell(22, 3, "=SUM(C7:C21)").number_format = PCT
    a.cell(22, 5, "=SUM(E7:E21)").number_format = PCT
    a.cell(22, 6, "=SUM(F7:F21)").number_format = MON0
    fila_total(a, 22, 1, 6)
    a.cell(24, 1, "Seguimiento semanal").font = Font(name=F, bold=True, color=ACENTO)
    encabezados(a, 25, [("Semana", 34), ("Fecha de corte", 15), ("% programado acumulado", 8), ("% real acumulado", 10), ("Desviación", 10),
                        ("Valor planeado PV", 15), ("", 3), ("Valor ganado EV", 14), ("Costo real AC", 14), ("SPI", 10), ("CPI", 10)])
    a.cell(25, 7).fill = PatternFill(None)
    prog = [0.03, 0.10, 0.21, 0.35, 0.50, 0.65, 0.78, 0.88, 0.95, 1.0]
    real = [0.02, 0.08, 0.18, 0.31]
    ini = fecha_inicio(D)
    for i in range(30):
        r = 26 + i
        celda(a, r, 1, i + 1, "0")
        celda(a, r, 2, ini + timedelta(days=7 * i + 5), FECHA, entrada=True)
        celda(a, r, 3, prog[i] if i < len(prog) else (1.0 if i < 12 else None), PCT, entrada=True)
        celda(a, r, 4, real[i] if i < len(real) else None, PCT, entrada=True)
        celda(a, r, 5, f"=IF(D{r}=\"\",\"\",D{r}-C{r})", PCT)
        celda(a, r, 6, f"=IF(C{r}=\"\",\"\",C{r}*$B$22)", MON0)
        celda(a, r, 8, f"=IF(D{r}=\"\",\"\",D{r}*$B$22)", MON0)
        celda(a, r, 9, round(monto * real[i] * 0.83, 2) if i < len(real) else None, MON0, entrada=True)
        celda(a, r, 10, f"=IF(OR(D{r}=\"\",F{r}=\"\",F{r}=0),\"\",H{r}/F{r})", "0.00")
        celda(a, r, 11, f"=IF(OR(H{r}=\"\",I{r}=\"\",I{r}=0),\"\",H{r}/I{r})", "0.00")
    a.cell(57, 1, "SPI / CPI: ≥ 1.00 verde, 0.90–0.99 ámbar, < 0.90 rojo. SPI < 1 = vas atrasado; CPI < 1 = gastas más de lo que ganas. El costo real AC es costo, sin utilidad.").font = Font(name=F, size=8, italic=True)
    for col in ("J", "K"):
        rng = f"{col}26:{col}55"
        a.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({col}26),{col}26<0.9)'], fill=ROJO))
        a.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({col}26),{col}26>=0.9,{col}26<1)'], fill=AMBAR))
        a.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({col}26),{col}26>=1)'], fill=VERDE))
    ch = LineChart(); ch.title = "Curva S: programado vs real"; ch.height = 8; ch.width = 18
    ch.y_axis.number_format = "0%"; ch.y_axis.scaling.max = 1; ch.y_axis.scaling.min = 0
    ch.add_data(Reference(a, min_col=3, min_row=25, max_row=55), titles_from_data=True)
    ch.add_data(Reference(a, min_col=4, min_row=25, max_row=55), titles_from_data=True)
    ch.set_categories(Reference(a, min_col=1, min_row=26, max_row=55))
    a.add_chart(ch, "H6")

    rd = wb.create_sheet("Reporte diario"); membrete(rd, "Reporte diario de obra", D, 12)
    encabezados(rd, 6, [("Fecha", 11), ("Clima", 10), ("Personal propio", 9), ("Personal subcontratistas", 11), ("Equipo activo", 18),
                        ("Actividades realizadas (frente / eje / nivel)", 44), ("Cantidades ejecutadas", 22), ("Incidencias / retrasos y causa", 30),
                        ("Instrucciones recibidas", 26), ("Asiento de bitácora No.", 10), ("Fotos (carpeta)", 14), ("Elaboró", 14)])
    ej = [[ini, "Despejado", 8, 6, "Andamio (4 cuerpos)", "Trazo y bastidor de muros de tablaroca ejes B-D / 2", "38 m² de bastidor", "Llegó tarde el material de canal (2 h)", "Supervisión pide cambiar altura de plafón en baños a 2.70 m", 12, "RD-001", "Superintendente"],
          [ini + timedelta(days=1), "Lluvia ligera", 8, 6, "Andamio (4 cuerpos)", "Forrado de muros 1 cara; ranurado para salidas eléctricas", "52 m² forrado", "Ninguna", "", 13, "RD-002", "Superintendente"]]
    for i in range(200):
        r = 7 + i
        for j in range(12):
            v = ej[i][j] if i < len(ej) else None
            x = celda(rd, r, j + 1, v, FECHA if j == 0 else ("0" if j in (2, 3, 9) else None), entrada=True)
            x.alignment = Alignment(wrap_text=True, vertical="top")
    lista(rd, "B7:B206", ["Despejado", "Nublado", "Lluvia ligera", "Lluvia fuerte", "Viento fuerte", "Calor extremo"])
    rd.freeze_panes = "B7"

    oc = wb.create_sheet("Órdenes de cambio"); membrete(oc, "Registro de órdenes de cambio", D, 10)
    encabezados(oc, 6, [("No.", 7), ("Fecha", 11), ("Descripción del cambio", 44), ("Origen", 14), ("Importe (+/-) sin IVA", 15),
                        ("Días de plazo (+/-)", 9), ("Estatus", 12), ("Fecha de autorización", 12), ("Autorizó", 18), ("Asiento de bitácora", 10)])
    ej = [["OC-01", ini + timedelta(days=1), "Cambio de altura de plafón en baños de 3.00 a 2.70 m con registro", "Propietario", 4800, 0, "Autorizada", ini + timedelta(days=3), "Representante del propietario", 13],
          ["OC-02", ini + timedelta(days=4), "Cuatro contactos adicionales en área de cajas", "Proyecto", 3560, 1, "Propuesta", None, None, None]]
    for i in range(60):
        r = 7 + i
        for j in range(10):
            v = ej[i][j] if i < len(ej) else None
            fmt = FECHA if j in (1, 7) else (MON if j == 4 else ("0" if j == 5 else None))
            x = celda(oc, r, j + 1, v, fmt, entrada=True); x.alignment = Alignment(wrap_text=True, vertical="top")
    lista(oc, "D7:D66", ["Propietario", "Proyecto", "Obra", "Autoridad", "Contratista"])
    lista(oc, "G7:G66", ["Propuesta", "Autorizada", "Rechazada"])
    resumen = [("Importe autorizado", '=SUMIFS(E7:E66,G7:G66,"Autorizada")', MON), ("Importe en trámite", '=SUMIFS(E7:E66,G7:G66,"Propuesta")', MON),
               ("Días de prórroga autorizados", '=SUMIFS(F7:F66,G7:G66,"Autorizada")', "0"), ("Monto original del contrato", g(D, "contrato.monto", 0), MON),
               ("Monto vigente (original + autorizadas)", "=M10+M7", MON), ("Variación sobre el original", "=IF(M10=0,0,M7/M10)", PCT)]
    oc.column_dimensions["L"].width = 34; oc.column_dimensions["M"].width = 16
    for i, (et, fo, fmt) in enumerate(resumen, 7):
        oc.cell(i, 12, et).font = B_FONT
        celda(oc, i, 13, fo, fmt, entrada=(i == 10))
    oc.cell(13, 12, "Sin orden firmada no se ejecuta ni se cobra (cláusula de ÓRDENES DE CAMBIO; art. 2500 CC B.C. en precio alzado).").font = Font(name=F, size=8, italic=True)

    pe = wb.create_sheet("Pendientes y RFI"); membrete(pe, "Pendientes, minutas y consultas técnicas (RFI)", D, 10)
    encabezados(pe, 6, [("No.", 8), ("Fecha", 11), ("Tipo", 10), ("Tema / pregunta", 44), ("Responsable", 18), ("Fecha compromiso", 12),
                        ("Fecha de cierre", 12), ("Estatus", 11), ("Días vencido", 9), ("Respuesta / acuerdo", 40)])
    ej = [["RFI-01", ini, "RFI", "Confirmar modelo de luminaria en área de cajas", "Proyectista", ini + timedelta(days=3), ini + timedelta(days=2), None, None, "Se confirma luminaria 60×60 4000 K"],
          ["P-01", ini + timedelta(days=2), "Pendiente", "Entrega de planos de cocineta por el propietario", "Propietario", ini + timedelta(days=5), None, None, None, ""]]
    for i in range(150):
        r = 7 + i
        for j in range(10):
            if j == 7:
                celda(pe, r, 8, f'=IF(D{r}="","",IF(G{r}<>"","Cerrado",IF(AND(F{r}<>"",F{r}<TODAY()),"Vencido","Abierto")))')
            elif j == 8:
                celda(pe, r, 9, f'=IF(AND(D{r}<>"",G{r}="",F{r}<>"",F{r}<TODAY()),TODAY()-F{r},0)', "0")
            else:
                v = ej[i][j] if i < len(ej) else None
                x = celda(pe, r, j + 1, v, FECHA if j in (1, 5, 6) else None, entrada=True); x.alignment = Alignment(wrap_text=True, vertical="top")
    lista(pe, "C7:C156", ["RFI", "Pendiente", "Minuta", "Calidad", "Seguridad"])
    pe.conditional_formatting.add("H7:H156", FormulaRule(formula=['H7="Vencido"'], fill=ROJO))
    pe.conditional_formatting.add("H7:H156", FormulaRule(formula=['H7="Cerrado"'], fill=VERDE))

    li = wb.create_sheet("Liberaciones y pruebas"); membrete(li, "Liberaciones de calidad y pruebas de laboratorio", D, 12)
    encabezados(li, 6, [("Fecha", 11), ("Elemento", 22), ("Eje / nivel", 14), ("Tipo de liberación o prueba", 22), ("Especificación (p. ej. f'c kg/cm²)", 13),
                        ("Resultado obtenido", 12), ("% de cumplimiento", 10), ("¿Cumple?", 9), ("Laboratorio / informe", 16), ("Liberó", 16), ("Asiento bitácora", 9), ("Observaciones", 30)])
    ej = [[ini + timedelta(days=28), "Losa de entrepiso", "Ejes 1-3 / A-C, N1", "Resistencia de concreto 28 días", 250, 262, None, None, "Lab. ejemplo / INF-118", "DRO", 20, ""],
          [ini + timedelta(days=6), "Relleno compactado", "Rampa de acceso", "Compactación Proctor", 0.95, 0.93, None, None, "Lab. ejemplo / INF-097", "Supervisión", 15, "Recompactar y volver a probar"]]
    for i in range(150):
        r = 7 + i
        for j in range(12):
            if j == 6:
                celda(li, r, 7, f'=IF(OR(E{r}="",F{r}="",E{r}=0),"",F{r}/E{r})', PCT)
            elif j == 7:
                celda(li, r, 8, f'=IF(G{r}="","",IF(G{r}>=1,"Sí","No"))')
            else:
                v = ej[i][j] if i < len(ej) else None
                x = celda(li, r, j + 1, v, FECHA if j == 0 else None, entrada=True); x.alignment = Alignment(wrap_text=True, vertical="top")
    li.conditional_formatting.add("H7:H156", FormulaRule(formula=['H7="No"'], fill=ROJO))
    li.conditional_formatting.add("H7:H156", FormulaRule(formula=['H7="Sí"'], fill=VERDE))
    return guardar(wb, carpeta, "Control_de_Obra.xlsx")

# ======================================================================
# 4. CONTROL DE PERSONAL
# ======================================================================
def personal(D, carpeta):
    wb = Workbook()
    instrucciones(wb, "Control de personal, lista de raya y subcontratistas REPSE", [
        "Para qué sirve:",
        "Plantilla de personal en obra, asistencia semanal con horas extra, lista de raya con costo real (FSR) y control documental de subcontratistas REPSE para saber a quién retenerle el pago.",
        "Cómo se llena:",
        "1. «Datos»: lunes de la semana, jornada máxima semanal y el FSR. El FSR sale de VERTICE-CALCULO (costos.BaseDatos().fsr(sd)) con IMSS, INFONAVIT e ISN B.C. 2026.",
        "2. «Plantilla»: alta de cada trabajador antes de que entre a obra (alta IMSS previa, SIROC, DC-3 y EPP).",
        "3. «Asistencia»: horas trabajadas por día. Las horas extra se separan en dobles y triples con las reglas de «Datos».",
        "4. «Lista de raya»: captura destajos, retenciones de IMSS e ISR (de tu sistema de nómina) y préstamos; el neto y el costo se calculan solos.",
        "5. «Subcontratistas REPSE»: cada mes marca Sí/No por documento. Si falta uno o el REPSE vence, el estatus dice RETENER PAGO (cláusula de REGISTRO REPSE del subcontrato; art. 15 LFT).",
        "Jornada:",
        "Hoy la jornada diurna máxima es de 48 horas semanales y las primeras 9 horas extra a la semana se pagan dobles; el excedente, triple (arts. 61, 66-68 LFT). La reforma constitucional publicada en el DOF el 3-mar-2026 reduce la jornada a 46 h en 2027, 44 h en 2028, 42 h en 2029 y 40 h en 2030: cambia los valores de «Datos» cuando entre en vigor cada etapa y confirma las reglas de horas extra de la reforma a la LFT.",
        "Privacidad:",
        "Este libro tiene datos personales (NSS, CURP). Guárdalo con acceso restringido conforme a la Ley Federal de Protección de Datos Personales en Posesión de los Particulares.",
    ])
    d = wb.create_sheet("Datos"); membrete(d, "Control de personal — parámetros", D, 3)
    d.column_dimensions["A"].width = 44; d.column_dimensions["B"].width = 16; d.column_dimensions["C"].width = 70
    param(d, 6, "Semana que inicia (lunes)", fecha_inicio(D), FECHA)
    param(d, 7, "Jornada máxima semanal (horas)", 48, "0", nota="48 h en 2026; 46 h desde 2027 (reforma DOF 3-mar-2026)")
    param(d, 8, "Horas extra semanales pagadas al doble", 9, "0", nota="Art. 67-68 LFT vigente; ajusta con la reforma")
    param(d, 9, "Horas de la jornada diaria", 8, "0")
    param(d, 10, "Salario mínimo diario ZLFN 2026", 440.87, MON, nota="Zona Libre de la Frontera Norte (fuente: mano-obra.json de VERTICE-CALCULO)")
    param(d, 11, "Factor de salario real (FSR)", 1.62, "0.0000", nota="Calcúlalo con VERTICE-CALCULO; depende del SD y de tu prima de riesgo")

    pl = wb.create_sheet("Plantilla"); membrete(pl, "Plantilla de personal en obra", D, 13)
    cols = [("No.", 6), ("Nombre completo", 28), ("Categoría", 16), ("Empresa (propio / subcontratista)", 18), ("NSS", 13), ("CURP", 20),
            ("Fecha alta IMSS", 11), ("Salario diario", 11), ("Ingreso a obra", 11), ("DC-3 vigente", 8), ("EPP entregado (fecha)", 11), ("Activo", 7), ("Alerta", 22)]
    encabezados(pl, 6, cols)
    ini = fecha_inicio(D)
    ej = [("Maestro de obra (ejemplo)", "Maestro", 1050), ("Oficial albañil (ejemplo)", "Oficial albañil", 630), ("Oficial tablaroquero (ejemplo)", "Oficial yesero", 640),
          ("Ayudante general (ejemplo)", "Peón", 505), ("Oficial electricista (ejemplo)", "Oficial electricista", 710), ("Ayudante general 2 (ejemplo)", "Peón", 430)]
    for i in range(80):
        r = 7 + i
        e = ej[i] if i < len(ej) else None
        celda(pl, r, 1, i + 1, "0")
        celda(pl, r, 2, e[0] if e else None, entrada=True)
        celda(pl, r, 3, e[1] if e else None, entrada=True)
        celda(pl, r, 4, "Propio" if e else None, entrada=True)
        celda(pl, r, 5, None, entrada=True); celda(pl, r, 6, None, entrada=True)
        celda(pl, r, 7, ini - timedelta(days=1) if e else None, FECHA, entrada=True)
        celda(pl, r, 8, e[2] if e else None, MON, entrada=True)
        celda(pl, r, 9, ini if e else None, FECHA, entrada=True)
        celda(pl, r, 10, ("Sí" if i < 5 else "No") if e else None, entrada=True)
        celda(pl, r, 11, ini if e else None, FECHA, entrada=True)
        celda(pl, r, 12, "Sí" if e else None, entrada=True)
        celda(pl, r, 13, f'=IF(B{r}="","",IF(H{r}<Datos!$B$10,"SD menor al mínimo ZLFN",IF(OR(G{r}="",G{r}>I{r}),"Alta IMSS posterior al ingreso",IF(K{r}="","Sin EPP",""))))')
    lista(pl, "J7:J86", ["Sí", "No", "N/A"]); lista(pl, "L7:L86", ["Sí", "No"])
    pl.conditional_formatting.add("M7:M86", FormulaRule(formula=['LEN(M7)>0'], fill=ROJO))
    pl.freeze_panes = "C7"

    asi = wb.create_sheet("Asistencia"); membrete(asi, "Asistencia y horas de la semana", D, 15)
    dias = ["Lun", "Mar", "Mié", "Jue", "Vie", "Sáb", "Dom"]
    cols = [("No.", 6), ("Nombre", 28), ("Categoría", 16)] + [(x, 6) for x in dias] + [("Horas totales", 8), ("Días trabajados", 8), ("Horas ordinarias", 8), ("Extra dobles", 8), ("Extra triples", 8)]
    encabezados(asi, 6, cols)
    for j in range(7):
        asi.cell(5, 4 + j, f"=Datos!$B$6+{j}").number_format = "dd/mm"
    horas = [9, 9, 9, 9, 9, 6, 0]
    for i in range(80):
        r = 7 + i; pr = 7 + i
        celda(asi, r, 1, f"=Plantilla!A{pr}", "0")
        celda(asi, r, 2, f'=IF(Plantilla!B{pr}="","",Plantilla!B{pr})')
        celda(asi, r, 3, f'=IF(Plantilla!C{pr}="","",Plantilla!C{pr})')
        for j in range(7):
            celda(asi, r, 4 + j, (horas[j] if i < 6 else None), "0.0", entrada=True)
        celda(asi, r, 11, f"=SUM(D{r}:J{r})", "0.0")
        celda(asi, r, 12, f'=COUNTIF(D{r}:J{r},">0")', "0")
        celda(asi, r, 13, f"=MIN(K{r},Datos!$B$7)", "0.0")
        celda(asi, r, 14, f"=MIN(MAX(0,K{r}-Datos!$B$7),Datos!$B$8)", "0.0")
        celda(asi, r, 15, f"=MAX(0,K{r}-Datos!$B$7-Datos!$B$8)", "0.0")
    asi.conditional_formatting.add("O7:O86", FormulaRule(formula=['O7>0'], fill=AMBAR))
    asi.freeze_panes = "D7"

    lr = wb.create_sheet("Lista de raya"); membrete(lr, "Lista de raya semanal", D, 15)
    cols = [("No.", 6), ("Nombre", 28), ("Salario diario", 10), ("Días trabajados", 8), ("Sueldo", 11), ("Séptimo día", 10), ("Horas extra", 10),
            ("Destajo / bonos", 10), ("Total percepciones", 12), ("Retención IMSS", 10), ("Retención ISR", 10), ("Préstamos / anticipos", 10),
            ("NETO A PAGAR", 12), ("Costo real empresa", 12), ("Firma", 16)]
    encabezados(lr, 6, cols)
    for i in range(80):
        r = 7 + i; pr = 7 + i
        celda(lr, r, 1, f"=Plantilla!A{pr}", "0")
        celda(lr, r, 2, f"=Asistencia!B{r}")
        celda(lr, r, 3, f'=IF(Plantilla!H{pr}="",0,Plantilla!H{pr})', MON)
        celda(lr, r, 4, f"=MIN(6,Asistencia!L{r})", "0")
        celda(lr, r, 5, f"=C{r}*D{r}", MON)
        celda(lr, r, 6, f"=C{r}*D{r}/6", MON)
        celda(lr, r, 7, f"=C{r}/Datos!$B$9*(2*Asistencia!N{r}+3*Asistencia!O{r})", MON)
        celda(lr, r, 8, None, MON, entrada=True)
        celda(lr, r, 9, f"=SUM(E{r}:H{r})", MON)
        for c in (10, 11, 12):
            celda(lr, r, c, None, MON, entrada=True)
        celda(lr, r, 13, f"=I{r}-J{r}-K{r}-L{r}", MON)
        celda(lr, r, 14, f"=(E{r}+F{r})*Datos!$B$11+G{r}+H{r}", MON)
        celda(lr, r, 15, None)
    tr = 87
    lr.cell(tr, 2, "TOTAL DE LA SEMANA")
    for c in range(5, 15):
        lr.cell(tr, c, f"=SUM({L(c)}7:{L(c)}86)").number_format = MON
    fila_total(lr, tr, 1, 15)
    lr.cell(tr + 1, 2, "Séptimo día proporcional (días trabajados / 6). Horas extra = salario por hora × 2 o × 3. El costo real aplica el FSR a sueldo y séptimo día.").font = Font(name=F, size=8, italic=True)
    firmas(lr, tr + 4, ["Elaboró — administrador de obra", "Autorizó — superintendente"], 2, 6)
    lr.freeze_panes = "C7"

    sr = wb.create_sheet("Subcontratistas REPSE"); membrete(sr, "Control documental de subcontratistas (REPSE)", D, 19)
    cols = [("Empresa", 26), ("RFC", 14), ("Especialidad", 18), ("No. REPSE", 14), ("Vigencia REPSE", 11), ("Días para vencer", 8), ("Trabajadores en obra", 8),
            ("Vinculado a SIROC", 8), ("Mes revisado", 9), ("CFDI nómina", 8), ("Pago IMSS", 8), ("Pago INFONAVIT", 8), ("Decl. ISR retenciones", 8),
            ("Decl. IVA y acuse", 8), ("ICSOE (cuatrim.)", 8), ("SISUB (cuatrim.)", 8), ("Opinión 32-D positiva", 8), ("ESTATUS", 16), ("Observaciones", 24)]
    encabezados(sr, 6, cols)
    ej = [("Instalaciones Ejemplo, S.A. de C.V.", "", "Eléctrica", "REPSE-EJEMPLO-1", ini + timedelta(days=400), 6, "Sí", "Sí", "Sí", "Sí", "Sí", "Sí", "Sí", "Sí", "Sí"),
          ("Plafones Ejemplo, S. de R.L.", "", "Tablaroca y plafones", "REPSE-EJEMPLO-2", ini + timedelta(days=20), 8, "Sí", "Sí", "No", "Sí", "Sí", "Sí", "Sí", "Sí", "Sí")]
    for i in range(40):
        r = 7 + i
        e = ej[i] if i < len(ej) else None
        vals = [e[0], e[1], e[2], e[3], e[4], None, e[5], e[6], ini.strftime("%m/%Y"), *e[7:]] if e else [None] * 17
        for j in range(17):
            if j == 5:
                celda(sr, r, 6, f'=IF(E{r}="","",E{r}-TODAY())', "0")
            else:
                celda(sr, r, j + 1, vals[j], FECHA if j == 4 else None, entrada=True)
        celda(sr, r, 18, f'=IF(A{r}="","",IF(OR(D{r}="",AND(ISNUMBER(F{r}),F{r}<0),COUNTIF(H{r}:Q{r},"No")>0),"RETENER PAGO",IF(AND(ISNUMBER(F{r}),F{r}<=30),"Renovar REPSE","OK")))')
        celda(sr, r, 19, None, entrada=True)
    lista(sr, "H7:Q46", ["Sí", "No", "N/A"])
    sr.conditional_formatting.add("R7:R46", FormulaRule(formula=['R7="RETENER PAGO"'], fill=ROJO))
    sr.conditional_formatting.add("R7:R46", FormulaRule(formula=['R7="Renovar REPSE"'], fill=AMBAR))
    sr.conditional_formatting.add("R7:R46", FormulaRule(formula=['R7="OK"'], fill=VERDE))
    sr.cell(48, 1, "Base: arts. 13-15 LFT (REPSE y responsabilidad solidaria), 15-A LSS (ICSOE), 29 Bis Ley INFONAVIT (SISUB), 27 fr. V LISR, 5 fr. II LIVA, 15-D y 32-D CFF. Sin REPSE capturado el estatus es RETENER PAGO.").font = Font(name=F, size=8, italic=True)
    sr.freeze_panes = "B7"
    return guardar(wb, carpeta, "Control_de_Personal.xlsx")

# ======================================================================
# 5. CONTROL DE MAQUINARIA (PROPIA Y RENTADA)
# ======================================================================
def maquinaria(D, carpeta):
    wb = Workbook()
    instrucciones(wb, "Control de maquinaria y equipo propio y rentado", [
        "Para qué sirve:",
        "Inventario de equipo propio y rentado, reporte diario de uso con horómetro, costo horario del equipo propio, conciliación de rentas contra la factura del arrendador y programa de mantenimiento por horómetro.",
        "Cómo se llena:",
        "1. «Inventario»: una fila por equipo con clave única. Para rentados captura proveedor, tarifa, unidad y mínimo (los del Anexo A del contrato de arrendamiento).",
        "2. «Uso diario»: una fila por equipo por día con horómetro inicial y final, firmada por el operador y tu representante (cláusula RENTA del contrato de arrendamiento). Anota horas inactivas y su causa: las horas por falla del equipo rentado no se pagan.",
        "3. «Costo horario propio»: método de costo horario (depreciación, inversión, seguros, mantenimiento, combustible, lubricantes, llantas y operación). El resultado alimenta el inventario y tus APU en VERTICE-CALCULO.",
        "4. «Control de rentas»: captura el periodo y lo que te facturó el arrendador; el libro calcula lo que debes pagar con el mínimo pactado y marca si la factura no cuadra.",
        "5. «Mantenimiento»: intervalos por horómetro; el estatus avisa cuándo programar el servicio.",
    ])
    ini = fecha_inicio(D)
    inv = wb.create_sheet("Inventario"); membrete(inv, "Inventario de maquinaria y equipo", D, 15)
    cols = [("Clave", 9), ("Equipo", 26), ("Tipo", 9), ("Marca", 12), ("Modelo", 12), ("Serie", 14), ("Proveedor (si es rentado)", 22), ("Tarifa (sin IVA)", 11),
            ("Unidad de tarifa", 9), ("Mínimo por día (h o jornadas)", 10), ("Costo horario propio", 11), ("Ingreso a obra", 11), ("Salida de obra", 11), ("Horómetro actual", 10), ("Estatus", 12)]
    encabezados(inv, 6, cols)
    eqs = list(g(D, "arrendamiento.equipos", []))
    filas = [("EQ-01", "Revolvedora de 1 saco", "Propio", "", "", "", "", None, "hora", None),
             ("EQ-02", "Vibrador de chicote a gasolina", "Propio", "", "", "", "", None, "hora", None)]
    for k, e in enumerate(eqs, 3):
        filas.append((f"EQ-0{k}", e.get("descripcion", ""), "Rentado", e.get("marca", ""), e.get("modelo", ""), e.get("serie", ""),
                      g(D, "contraparte.nombre") or "Arrendador", e.get("tarifa"), e.get("unidad", "hora"), 8 if e.get("unidad", "hora") == "hora" else 1))
    for i in range(40):
        r = 7 + i
        e = filas[i] if i < len(filas) else None
        for j in range(10):
            celda(inv, r, j + 1, e[j] if e else None, MON if j == 7 else ("0.0" if j == 9 else None), entrada=True)
        celda(inv, r, 11, f'=IF(C{r}<>"Propio","",IFERROR(INDEX(\'Costo horario propio\'!$B$29:$F$29,MATCH(A{r},\'Costo horario propio\'!$B$6:$F$6,0)),"Calcula en hoja"))', MON)
        celda(inv, r, 12, ini if e else None, FECHA, entrada=True)
        celda(inv, r, 13, None, FECHA, entrada=True)
        celda(inv, r, 14, f'=IF(A{r}="","",_xlfn.MAXIFS(\'Uso diario\'!$F$7:$F$1006,\'Uso diario\'!$B$7:$B$1006,A{r}))', "#,##0.0")
        celda(inv, r, 15, f'=IF(A{r}="","",IF(M{r}<>"","Fuera de obra","En obra"))')
    lista(inv, "C7:C46", ["Propio", "Rentado"]); lista(inv, "I7:I46", ["hora", "jornada", "semana", "mes"])
    inv.freeze_panes = "C7"

    uso = wb.create_sheet("Uso diario"); membrete(uso, "Reporte diario de uso de equipo", D, 14)
    cols = [("Fecha", 11), ("Clave", 9), ("Equipo", 24), ("Tipo", 9), ("Horómetro inicial", 10), ("Horómetro final", 10), ("Horas trabajadas", 9), ("Horas inactivas", 9),
            ("Causa de inactividad", 20), ("Combustible (L)", 9), ("Rendimiento (L/h)", 9), ("Frente / concepto", 24), ("Operador", 16), ("Firma de conformidad", 14)]
    encabezados(uso, 6, cols)
    ej = [(ini, "EQ-03", 1250.0, 1258.5, 0, "", 72, "Excavación de cepas eje A"), (ini, "EQ-04", None, None, 0, "", 4, "Compactación de relleno"),
          (ini + timedelta(days=1), "EQ-03", 1258.5, 1264.0, 2.5, "Falla hidráulica (equipo rentado)", 48, "Excavación de cepas eje B"),
          (ini + timedelta(days=1), "EQ-01", None, None, 0, "", 6, "Colado de firmes")]
    for i in range(1000):
        r = 7 + i
        e = ej[i] if i < len(ej) else None
        celda(uso, r, 1, e[0] if e else None, FECHA, entrada=True)
        celda(uso, r, 2, e[1] if e else None, entrada=True)
        celda(uso, r, 3, f'=IF(B{r}="","",IFERROR(INDEX(Inventario!$B$7:$B$46,MATCH(B{r},Inventario!$A$7:$A$46,0)),"Clave no existe"))')
        celda(uso, r, 4, f'=IF(B{r}="","",IFERROR(INDEX(Inventario!$C$7:$C$46,MATCH(B{r},Inventario!$A$7:$A$46,0)),""))')
        celda(uso, r, 5, e[2] if e else None, "#,##0.0", entrada=True)
        celda(uso, r, 6, e[3] if e else None, "#,##0.0", entrada=True)
        celda(uso, r, 7, f'=IF(OR(E{r}="",F{r}=""),IF(B{r}="","",IF(A{r}="","",8)),F{r}-E{r})', "0.0")
        uso.cell(r, 7).comment = None
        celda(uso, r, 8, e[4] if e else None, "0.0", entrada=True)
        celda(uso, r, 9, e[5] if e else None, entrada=True)
        celda(uso, r, 10, e[6] if e else None, "0.0", entrada=True)
        celda(uso, r, 11, f'=IF(OR(J{r}="",N(G{r})=0),"",J{r}/G{r})', "0.0")
        celda(uso, r, 12, e[7] if e else None, entrada=True)
        celda(uso, r, 13, None, entrada=True); celda(uso, r, 14, None, entrada=True)
    uso.cell(6, 7).comment = Comment("Horómetro final − inicial. Si el equipo no tiene horómetro se toman 8 h por día reportado; corrige la fórmula si trabajó menos.", "Vértice")
    uso.freeze_panes = "C7"

    ch = wb.create_sheet("Costo horario propio"); membrete(ch, "Costo horario de equipo propio", D, 6)
    ch.column_dimensions["A"].width = 46
    for col in "BCDEF":
        ch.column_dimensions[col].width = 15
    for j, (clave, nom) in enumerate([("EQ-01", "Revolvedora 1 saco"), ("EQ-02", "Vibrador gasolina"), ("", ""), ("", ""), ("", "")]):
        c1 = celda(ch, 6, 2 + j, clave or None, entrada=True); c1.font = Font(name=F, bold=True, color="0000FF", size=9)
        celda(ch, 7, 2 + j, nom or None, entrada=True)
    ch.cell(6, 1, "Clave del equipo (igual al inventario)").font = B_FONT
    ch.cell(7, 1, "Equipo").font = B_FONT
    datos = [  # etiqueta, ej1, ej2, formato
        ("Valor de adquisición Va (sin IVA)", 38000, 14500, MON), ("Valor de llantas Pn (si tiene)", 0, 0, MON),
        ("Valor de rescate (% de Va)", 0.15, 0.10, PCT), ("Vida económica Ve (horas)", 6000, 3000, "#,##0"),
        ("Horas trabajadas por año Hea", 1200, 1000, "#,##0"), ("Tasa de interés anual i", 0.14, 0.14, PCT),
        ("Prima anual de seguros s", 0.03, 0.03, PCT), ("Coeficiente de mantenimiento Ko (× depreciación)", 0.8, 0.9, "0.00"),
        ("Consumo de combustible (L/h)", 0.9, 0.6, "0.00"), ("Precio del combustible ($/L sin IVA)", 22.5, 22.5, MON),
        ("Lubricantes (% del costo de combustible)", 0.15, 0.20, PCT), ("Vida de llantas (horas)", 2000, 2000, "#,##0"),
        ("Salario real del operador por turno (con FSR)", 0, 0, MON), ("Horas efectivas por turno", 8, 8, "0.0"),
    ]
    for i, (et, v1, v2, fmt) in enumerate(datos, 8):
        ch.cell(i, 1, et).font = N_FONT
        for j in range(5):
            celda(ch, i, 2 + j, (v1, v2)[j] if j < 2 else None, fmt, entrada=True)
    calc = [
        ("Valor de rescate Vr", "={c}8*{c}10", MON), ("Depreciación D = (Va − Vr) / Ve", "=IF(N({c}11)=0,0,({c}8-{c}22)/{c}11)", MON),
        ("Inversión Im = (Va + Vr) / (2·Hea) · i", "=IF(N({c}12)=0,0,({c}8+{c}22)/(2*{c}12)*{c}13)", MON),
        ("Seguros Sm = (Va + Vr) / (2·Hea) · s", "=IF(N({c}12)=0,0,({c}8+{c}22)/(2*{c}12)*{c}14)", MON),
        ("Mantenimiento Mn = Ko · D", "={c}15*{c}23", MON),
        ("Combustible y lubricantes", "={c}16*{c}17*(1+{c}18)", MON),
        ("Llantas + operación", "=IF(N({c}19)=0,0,{c}9/{c}19)+IF(N({c}21)=0,0,{c}20/{c}21)", MON),
        ("COSTO HORARIO DIRECTO", "=SUM({c}23:{c}28)", MON),
    ]
    for i, (et, fo, fmt) in enumerate(calc, 22):
        ch.cell(i, 1, et).font = B_FONT if i == 29 else N_FONT
        for j in range(5):
            c = L(2 + j)
            celda(ch, i, 2 + j, f'=IF({c}$6="","",{fo.format(c=c)[1:]})', fmt, negrita=(i == 29))
    fila_total(ch, 29, 1, 6)
    ch.cell(31, 1, "Método de costo horario (cargos fijos, consumos y operación), el mismo que usan los análisis de precios unitarios. Si el equipo se usa en APU de VERTICE-CALCULO, captura este costo como Insumo tipo «Equipo».").font = Font(name=F, size=8, italic=True)

    rt = wb.create_sheet("Control de rentas"); membrete(rt, "Conciliación de rentas de equipo", D, 15)
    cols = [("Clave", 9), ("Equipo", 24), ("Proveedor", 20), ("Periodo del", 11), ("Periodo al", 11), ("Unidad de tarifa", 9), ("Días con uso", 8), ("Horas trabajadas", 9),
            ("Cantidad mínima del periodo", 10), ("Cantidad a pagar", 9), ("Tarifa", 10), ("Fletes", 10), ("IMPORTE CORRECTO", 13), ("Facturado por el proveedor", 13), ("Diferencia / revisión", 18)]
    encabezados(rt, 6, cols)
    rentados = [f for f in filas if f[2] == "Rentado"]
    for i in range(30):
        r = 7 + i
        e = rentados[i] if i < len(rentados) else None
        celda(rt, r, 1, e[0] if e else None, entrada=True)
        celda(rt, r, 2, f'=IF(A{r}="","",INDEX(Inventario!$B$7:$B$46,MATCH(A{r},Inventario!$A$7:$A$46,0)))')
        celda(rt, r, 3, f'=IF(A{r}="","",INDEX(Inventario!$G$7:$G$46,MATCH(A{r},Inventario!$A$7:$A$46,0)))')
        celda(rt, r, 4, ini if e else None, FECHA, entrada=True)
        celda(rt, r, 5, ini + timedelta(days=13) if e else None, FECHA, entrada=True)
        celda(rt, r, 6, f'=IF(A{r}="","",INDEX(Inventario!$I$7:$I$46,MATCH(A{r},Inventario!$A$7:$A$46,0)))')
        celda(rt, r, 7, f"=IF(A{r}=\"\",0,SUMPRODUCT(('Uso diario'!$B$7:$B$1006=A{r})*('Uso diario'!$A$7:$A$1006>=D{r})*('Uso diario'!$A$7:$A$1006<=E{r})*('Uso diario'!$G$7:$G$1006>0)))", "0")
        celda(rt, r, 8, f"=IF(A{r}=\"\",0,SUMIFS('Uso diario'!$G$7:$G$1006,'Uso diario'!$B$7:$B$1006,A{r},'Uso diario'!$A$7:$A$1006,\">=\"&D{r},'Uso diario'!$A$7:$A$1006,\"<=\"&E{r}))", "0.0")
        celda(rt, r, 9, f'=IF(A{r}="","",G{r}*INDEX(Inventario!$J$7:$J$46,MATCH(A{r},Inventario!$A$7:$A$46,0)))', "0.0")
        celda(rt, r, 10, f'=IF(A{r}="","",IF(F{r}="hora",MAX(H{r},I{r}),MAX(G{r},I{r})))', "0.0")
        celda(rt, r, 11, f'=IF(A{r}="","",INDEX(Inventario!$H$7:$H$46,MATCH(A{r},Inventario!$A$7:$A$46,0)))', MON)
        celda(rt, r, 12, float(g(D, "arrendamiento.flete", 0)) * 2 if (e and i == 0) else None, MON, entrada=True)
        celda(rt, r, 13, f'=IF(A{r}="","",J{r}*K{r}+N(L{r}))', MON)
        celda(rt, r, 14, None, MON, entrada=True)
        celda(rt, r, 15, f'=IF(OR(A{r}="",N{r}=""),"",IF(ABS(N{r}-M{r})<1,"Cuadra",IF(N{r}>M{r},"Cobro de más: $"&TEXT(N{r}-M{r},"#,##0.00"),"Cobro de menos: $"&TEXT(M{r}-N{r},"#,##0.00"))))')
    rt.conditional_formatting.add("O7:O36", FormulaRule(formula=['LEFT(O7,5)="Cobro"'], fill=ROJO))
    rt.cell(38, 1, "Las horas por falla del equipo rentado van en «Horas inactivas» y no se suman a las horas trabajadas, por lo que no se pagan (cláusula de MANTENIMIENTO del contrato de arrendamiento). Los «Días con uso» cuentan días con horas trabajadas mayores a cero. Para tarifas semanales o mensuales ajusta la cantidad a pagar con la factura.").font = Font(name=F, size=8, italic=True)

    mt = wb.create_sheet("Mantenimiento"); membrete(mt, "Programa de mantenimiento por horómetro", D, 9)
    cols = [("Clave", 9), ("Equipo", 24), ("Servicio", 26), ("Intervalo (horas)", 9), ("Horómetro del último servicio", 11), ("Próximo servicio (horómetro)", 11),
            ("Horómetro actual", 10), ("Horas restantes", 9), ("Estatus", 14)]
    encabezados(mt, 6, cols)
    ej = [("EQ-03", "Cambio de aceite de motor y filtros", 250, 1100), ("EQ-03", "Engrase general", 50, 1250), ("EQ-01", "Revisión de banda y tambor", 100, 0)]
    for i in range(60):
        r = 7 + i
        e = ej[i] if i < len(ej) else None
        celda(mt, r, 1, e[0] if e else None, entrada=True)
        celda(mt, r, 2, f'=IF(A{r}="","",IFERROR(INDEX(Inventario!$B$7:$B$46,MATCH(A{r},Inventario!$A$7:$A$46,0)),""))')
        celda(mt, r, 3, e[1] if e else None, entrada=True)
        celda(mt, r, 4, e[2] if e else None, "#,##0", entrada=True)
        celda(mt, r, 5, e[3] if e else None, "#,##0.0", entrada=True)
        celda(mt, r, 6, f'=IF(A{r}="","",E{r}+D{r})', "#,##0.0")
        celda(mt, r, 7, f'=IF(A{r}="","",IFERROR(INDEX(Inventario!$N$7:$N$46,MATCH(A{r},Inventario!$A$7:$A$46,0)),0))', "#,##0.0")
        celda(mt, r, 8, f'=IF(A{r}="","",F{r}-G{r})', "#,##0.0")
        celda(mt, r, 9, f'=IF(A{r}="","",IF(N(G{r})=0,"Sin lectura",IF(H{r}<=0,"VENCIDO",IF(H{r}<=0.1*D{r},"Programar","OK"))))')
    mt.conditional_formatting.add("I7:I66", FormulaRule(formula=['I7="VENCIDO"'], fill=ROJO))
    mt.conditional_formatting.add("I7:I66", FormulaRule(formula=['I7="Programar"'], fill=AMBAR))
    mt.conditional_formatting.add("I7:I66", FormulaRule(formula=['I7="OK"'], fill=VERDE))
    return guardar(wb, carpeta, "Control_de_Maquinaria.xlsx")

# ======================================================================
# 6. COTIZACIÓN
# ======================================================================
def cotizacion(D, carpeta):
    wb = Workbook()
    instrucciones(wb, "Cotización con análisis de precios", [
        "Para qué sirve:",
        "Arma la cotización al cliente a partir del costo directo de cada concepto y del factor de sobrecosto (indirectos, financiamiento, utilidad y cargos adicionales), con la misma fórmula de VERTICE-CALCULO (costos.Sobrecosto.factor).",
        "Cómo se llena:",
        "1. «Datos»: cliente, obra, vigencia, porcentajes de sobrecosto y condiciones comerciales.",
        "2. «Indirectos»: los gastos de campo de la obra (superintendente, residente, seguridad, renta de oficina, fianzas). Su porcentaje pasa solo a «Datos».",
        "3. «Análisis»: costo directo unitario de cada concepto por tipo de insumo. Toma los costos de los APU de VERTICE-CALCULO (precios de precios-obra-, FSR 2026 Tijuana) o de la hoja «APU».",
        "4. «Cotización»: es la hoja que se imprime o exporta a PDF para el cliente. No muestra costos ni márgenes.",
        "Cuidado:",
        "Una cotización aceptada se formaliza con el contrato de obra a precio alzado o a precios unitarios del skill; las condiciones de anticipo, pago y plazo de la cotización deben ser las mismas del contrato.",
    ])
    ini = fecha_inicio(D)
    d = wb.create_sheet("Datos"); membrete(d, "Cotización — datos y sobrecosto", D, 3)
    d.column_dimensions["A"].width = 40; d.column_dimensions["B"].width = 34; d.column_dimensions["C"].width = 50
    filas = [("Cotización número", "COT-2026-001", None), ("Fecha", date.today(), FECHA), ("Vigencia (días naturales)", 15, "0"),
             ("Cliente", g(D, "contraparte.nombre"), None), ("Atención", g(D, "contraparte.representante"), None), ("Obra", g(D, "obra.nombre"), None),
             ("Ubicación", g(D, "obra.ubicacion") or g(D, "obra.municipio"), None)]
    for i, (et, v, fmt) in enumerate(filas, 6):
        param(d, i, et, v, fmt)
    d.cell(14, 1, "SOBRECOSTO").font = Font(name=F, bold=True, color=ACENTO)
    param(d, 15, "Indirectos de oficina", 0.06, PCT)
    param(d, 16, "Indirectos de campo", "=Indirectos!E20", PCT, entrada=False, nota="Calculado en la hoja Indirectos")
    param(d, 17, "Financiamiento", 0.01, PCT, nota="Usa el costo financiero del libro Flujo_de_Caja")
    param(d, 18, "Utilidad", 0.10, PCT)
    param(d, 19, "Cargos adicionales", 0.0, PCT)
    param(d, 20, "Factor de sobrecosto", "=(1+B15+B16)*(1+B17)*(1+B18)*(1+B19)", "0.0000", entrada=False, nota="(1+ind)(1+fin)(1+util)(1+cargos), igual que VERTICE-CALCULO")
    param(d, 21, "IVA", float(g(D, "contrato.iva_pct", 8)) / 100, PCT, nota="8 % región fronteriza norte (Decreto vigente hasta 31-dic-2026)")
    d.cell(23, 1, "CONDICIONES COMERCIALES").font = Font(name=F, bold=True, color=ACENTO)
    cond = [("Anticipo", f"{g(D, 'contrato.anticipo_pct', 30)} % a la firma del contrato"),
            ("Forma de pago", f"Estimaciones {g(D, 'contrato.periodo_estimacion', 'quincenal')}es por avance, pagaderas a {g(D, 'contrato.dias_pago', 15)} días"),
            ("Tiempo de ejecución", f"{g(D, 'contrato.plazo_dias', 60)} días naturales a partir del anticipo y la entrega del sitio"),
            ("Incluye", "Materiales, mano de obra, herramienta, equipo, fletes, limpieza y retiro de escombro"),
            ("No incluye", "Licencias y derechos municipales, DRO y corresponsables, proyecto ejecutivo, mobiliario, trabajos no descritos"),
            ("Garantía", f"{g(D, 'contrato.vicios_meses', 12)} meses contra vicios ocultos")]
    for i, (et, v) in enumerate(cond, 24):
        param(d, i, et, v)

    ind = wb.create_sheet("Indirectos"); membrete(ind, "Indirectos de campo", D, 5)
    encabezados(ind, 6, [("Concepto", 40), ("Costo mensual / unitario", 15), ("Cantidad (meses o pzas)", 12), ("% de dedicación", 11), ("Total", 15)])
    meses = round(float(g(D, "contrato.plazo_dias", 60)) / 30, 1)
    ej = [("Superintendente de obra", 45000, meses, 0.10), ("Residente / auxiliar técnico", 28000, meses, 0.20), ("Responsable de seguridad (NOM-031)", 22000, meses, 0.10),
          ("Bodeguero / velador", 14000, meses, 0.20), ("Oficina, sanitario y bodega provisionales", 2500, meses, 1), ("Fianzas (anticipo, cumplimiento y vicios)", 6000, 1, 1),
          ("Seguro de responsabilidad civil", 4000, 1, 1), ("Pruebas de laboratorio", 2000, 1, 1), ("Limpieza final y señalización", 3000, 1, 1)]
    for i in range(12):
        r = 7 + i
        e = ej[i] if i < len(ej) else None
        celda(ind, r, 1, e[0] if e else None, entrada=True)
        celda(ind, r, 2, e[1] if e else None, MON, entrada=True)
        celda(ind, r, 3, e[2] if e else None, "0.0", entrada=True)
        celda(ind, r, 4, e[3] if e else None, PCT, entrada=True)
        celda(ind, r, 5, f"=N(B{r})*N(C{r})*N(D{r})", MON)
    ind.cell(19, 1, "Total de indirectos de campo"); ind.cell(19, 5, "=SUM(E7:E18)").number_format = MON; fila_total(ind, 19, 1, 5)
    ind.cell(20, 1, "Porcentaje sobre costo directo"); ind.cell(20, 5, "=IF(Análisis!J49=0,0,E19/Análisis!J49)").number_format = PCT; fila_total(ind, 20, 1, 5)

    an = wb.create_sheet("Análisis"); membrete(an, "Análisis de costo directo y precio unitario", D, 12)
    cols = [("Clave", 9), ("Concepto", 44), ("Unidad", 7), ("Cantidad", 10), ("Materiales", 11), ("Mano de obra", 11), ("Equipo", 10),
            ("Herramienta menor (3 % MO)", 10), ("Subcontrato", 11), ("Costo directo unitario", 12), ("P.U. de venta", 12), ("Importe de venta", 14)]
    encabezados(an, 6, cols)
    costos = [(70, 55, 10, 0), (210, 165, 8, 0), (250, 190, 10, 0), (380, 290, 0, 0), (930, 300, 0, 0), (1050, 700, 0, 0), (360, 200, 0, 0), (32, 38, 0, 0)]
    for i in range(42):
        r = 7 + i
        e = CONCEPTOS_EJ[i] if i < len(CONCEPTOS_EJ) else None
        cst = costos[i] if i < len(costos) else None
        celda(an, r, 1, e[0] if e else None, entrada=True)
        x = celda(an, r, 2, e[1] if e else None, entrada=True); x.alignment = Alignment(wrap_text=True)
        celda(an, r, 3, e[2] if e else None, entrada=True)
        celda(an, r, 4, e[3] if e else None, NUM, entrada=True)
        for j in range(3):
            celda(an, r, 5 + j, cst[j] if cst else None, MON, entrada=True)
        celda(an, r, 8, f"=N(F{r})*0.03", MON)
        celda(an, r, 9, cst[3] if cst else None, MON, entrada=True)
        celda(an, r, 10, f"=N(E{r})+N(F{r})+N(G{r})+H{r}+N(I{r})", MON)
        celda(an, r, 11, f"=ROUND(J{r}*Datos!$B$20,2)", MON)
        celda(an, r, 12, f"=K{r}*N(D{r})", MON)
    an.cell(49, 2, "TOTALES (costo directo y venta)")
    an.cell(49, 10, "=SUMPRODUCT(J7:J48,D7:D48)").number_format = MON
    an.cell(49, 12, "=SUM(L7:L48)").number_format = MON
    an.cell(50, 2, "Margen bruto sobre venta")
    an.cell(50, 12, "=IF(L49=0,0,(L49-J49)/L49)").number_format = PCT
    fila_total(an, 49, 1, 12); fila_total(an, 50, 1, 12)
    an.freeze_panes = "C7"

    ap = wb.create_sheet("APU"); membrete(ap, "Análisis de precio unitario (plantilla)", D, 6)
    for col, w in zip("ABCDEF", (40, 9, 12, 13, 13, 30)):
        ap.column_dimensions[col].width = w
    param(ap, 5, "Concepto", CONCEPTOS_EJ[1][1]); ap.cell(5, 3, "Unidad").font = B_FONT; celda(ap, 5, 4, "m²", entrada=True)
    secciones = [("MATERIALES", ["Insumo", "Unidad", "Cantidad", "Precio (sin IVA)", "Importe", "Fuente / cotización"],
                  [("Panel de yeso 1/2\" 1.22×2.44", "m²", 1.05, 112.0, "precios-obra- (ago 2026) + 10 % Tijuana"), ("Poste metálico cal. 26 de 3 m", "pza", 0.70, 68.0, "cotización"),
                   ("Canal metálico cal. 26 de 3 m", "pza", 0.30, 55.0, "cotización"), ("Tornillo, cinta y compuesto", "lote", 1, 22.0, "cotización")]),
                 ("MANO DE OBRA", ["Cuadrilla", "Unidad", "Cantidad (1/rend.)", "Salario real", "Importe", "Rendimiento (m²/jor)"],
                  [("1 oficial yesero + 1 ayudante (salario real con FSR)", "jor", "=1/F{r}", 1850.0, 11)]),
                 ("EQUIPO Y HERRAMIENTA", ["Equipo", "Unidad", "Cantidad", "Costo", "Importe", "Nota"],
                  [("Andamio tubular (cuerpo)", "jornada", 0.09, 35.0, "mano-obra.json"), ("Herramienta menor (% de MO)", "%", 0.03, "=E18", "3 % de la mano de obra")])]
    r = 7
    sumas = []
    for titulo, cab, filas_ in secciones:
        ap.cell(r, 1, titulo).font = Font(name=F, bold=True, color=ACENTO); r += 1
        encabezados(ap, r, [(c, None) for c in cab]); r += 1
        ini_r = r
        for fl in filas_:
            celda(ap, r, 1, fl[0], entrada=True); celda(ap, r, 2, fl[1], entrada=True)
            qty = fl[2].format(r=r) if isinstance(fl[2], str) else fl[2]
            celda(ap, r, 3, qty, "0.0000", entrada=not isinstance(fl[2], str))
            celda(ap, r, 4, fl[3], MON, entrada=not (isinstance(fl[3], str)))
            celda(ap, r, 5, f"=C{r}*D{r}", MON)
            celda(ap, r, 6, fl[4], "0.0" if isinstance(fl[4], (int, float)) else None, entrada=True)
            r += 1
        ap.cell(r, 4, "Subtotal").font = B_FONT
        ap.cell(r, 5, f"=SUM(E{ini_r}:E{r-1})").number_format = MON; fila_total(ap, r, 4, 5)
        sumas.append(r); r += 2
    ap.cell(r, 4, "COSTO DIRECTO").font = B_FONT
    ap.cell(r, 5, "=" + "+".join(f"E{s}" for s in sumas)).number_format = MON
    ap.cell(r + 1, 4, "× Factor de sobrecosto").font = N_FONT
    ap.cell(r + 1, 5, "=Datos!B20").number_format = "0.0000"
    ap.cell(r + 2, 4, "PRECIO UNITARIO").font = B_FONT
    ap.cell(r + 2, 5, f"=ROUND(E{r}*E{r+1},2)").number_format = MON
    fila_total(ap, r, 4, 5); fila_total(ap, r + 2, 4, 5)
    ap.cell(r + 4, 1, "Para APU con base de precios y FSR automáticos usa VERTICE-CALCULO: costos.apu_desde_base(db, rendimiento). Esta hoja es para análisis rápidos en obra.").font = Font(name=F, size=8, italic=True)

    ct = wb.create_sheet("Cotización"); membrete(ct, "Cotización", D, 6)
    for col, w in zip("ABCDEF", (10, 52, 8, 11, 14, 16)):
        ct.column_dimensions[col].width = w
    ct.cell(3, 1, "Cotización:").font = B_FONT; ct.cell(3, 2, "=Datos!B6")
    ct.cell(4, 1, "Contrato:").value = "Fecha:"; ct.cell(4, 2, "=Datos!B7").number_format = FECHA
    info = [("Cliente:", '=""&Datos!B9'), ("Atención:", '=""&Datos!B10'), ("Obra:", '=""&Datos!B11'), ("Ubicación:", '=""&Datos!B12')]
    for i, (et, fo) in enumerate(info, 5):
        ct.cell(i, 1, et).font = B_FONT; ct.cell(i, 2, fo).font = N_FONT
    encabezados(ct, 10, [("Clave", None), ("Descripción", None), ("Unidad", None), ("Cantidad", None), ("P.U. (sin IVA)", None), ("Importe", None)])
    for i in range(42):
        r = 11 + i; ar = 7 + i
        for j, col in enumerate("ABCD"):
            celda(ct, r, j + 1, f'=IF(Análisis!{col}{ar}="","",Análisis!{col}{ar})', NUM if col == "D" else None)
        ct.cell(r, 2).alignment = Alignment(wrap_text=True)
        celda(ct, r, 5, f'=IF(Análisis!A{ar}="","",Análisis!K{ar})', MON)
        celda(ct, r, 6, f'=IF(Análisis!A{ar}="","",Análisis!L{ar})', MON)
    t0 = 54
    for i, (et, fo) in enumerate([("Subtotal", "=SUM(F11:F52)"), ("IVA", f"=ROUND(F{t0}*Datos!B21,2)"), ("TOTAL", f"=F{t0}+F{t0+1}")]):
        ct.cell(t0 + i, 5, et).font = B_FONT
        ct.cell(t0 + i, 6, fo).number_format = MON
        fila_total(ct, t0 + i, 5, 6)
    ct.cell(58, 1, "Condiciones").font = Font(name=F, bold=True, color=ACENTO)
    for i in range(6):
        ct.cell(59 + i, 1, f"=Datos!A{24+i}").font = B_FONT
        ct.cell(59 + i, 2, f"=Datos!B{24+i}").font = N_FONT
    ct.cell(65, 1, "Vigencia:").font = B_FONT
    ct.cell(65, 2, '=Datos!B8&" días naturales a partir de la fecha de esta cotización. Precios en moneda nacional, sin IVA salvo donde se indica."').font = N_FONT
    ct.cell(66, 2, "La aceptación de esta cotización se formaliza mediante el contrato de obra correspondiente.").font = Font(name=F, size=8, italic=True)
    ct.cell(69, 2, "_" * 34).font = N_FONT
    ct.cell(70, 2, g(D, "empresa.titular") or "Ing. Miguel Estrada").font = B_FONT
    ct.cell(71, 2, g(D, "empresa.nombre_comercial")).font = N_FONT
    ct.cell(72, 2, " · ".join(x for x in [g(D, "empresa.telefono"), g(D, "empresa.correo")] if x) or "Teléfono · correo").font = N_FONT
    ct.print_area = "A1:F72"
    ct.page_setup.orientation = "portrait"; ct.page_setup.fitToWidth = 1; ct.page_setup.fitToHeight = 0
    ct.sheet_properties.pageSetUpPr.fitToPage = True
    wb.move_sheet("Cotización", offset=-4)
    return guardar(wb, carpeta, "Cotizacion.xlsx")

FORMATOS = {"flujo": flujo, "estimaciones": estimaciones, "obra": control_obra, "personal": personal, "maquinaria": maquinaria, "cotizacion": cotizacion}

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__); sys.exit(1)
    with open(sys.argv[1], encoding="utf-8") as fh:
        D = json.load(fh)
    os.makedirs(sys.argv[2], exist_ok=True)
    for nombre in (sys.argv[3:] or list(FORMATOS)):
        FORMATOS[nombre](D, sys.argv[2])
